"""Baseline V1：重新只读导入Excel、固定参数、配比簇随机划分与论文分组评价。"""
import csv
from datetime import datetime, timezone
import importlib.metadata
import json
from pathlib import Path
import shutil
from urllib.parse import unquote

import joblib
import numpy as np
import openpyxl
from sklearn.dummy import DummyRegressor
from sklearn.ensemble import RandomForestRegressor
from sklearn.impute import SimpleImputer
from sklearn.inspection import permutation_importance
from sklearn.model_selection import GroupShuffleSplit
from sklearn.pipeline import Pipeline

from .audit import identity_values
from .baselines import regression_metrics, grouped_folds
from .parallel_exploration import numeric, feature_matrix, digest, write_json

RECIPE_FIELDS=['Water','Cem','SF','FA','GGBS','MK','NP','AH','GP','QP','LS','RS','SS','CA','SCA',
               'SP','AWA','Retarder','Emulsion','StF','SyF','AEA','UPA','EA','AlkaliAct','w/c','w/b']


def inventory():
    """实际打开现存Excel，区分原始附件和项目派生报告，不假定缺失文件存在。"""
    paths=sorted(set(Path('/workspace/attachments').rglob('*.xlsx'))|set(Path.cwd().rglob('*.xlsx')))
    result=[]
    for p in paths:
        try:
            w=openpyxl.load_workbook(p,read_only=True,data_only=False)
            result.append({'path':str(p),'display_name':unquote(p.name),'kind':'original_attachment' if '/attachments/' in str(p) else 'derived_project_file',
                'sha256':digest(p),'sheets':[{'name':s.title,'rows':s.max_row,'columns':s.max_column} for s in w]})
            w.close()
        except Exception as exc:
            result.append({'path':str(p),'error':str(exc)})
    names={r.get('display_name') for r in result}
    missing=[n for n in ['水下混凝土_人工核查版.xlsx','水下混凝土_筛选结果_20261012.xlsx'] if n not in names]
    return {'workbooks':result,'missing_requested_workbooks':missing}


def read_source(path):
    """字段由参数说明表映射；仅展开原有身份合并单元格，避免普通空白串来源。"""
    path=Path(path);w=openpyxl.load_workbook(path,data_only=False)
    if 'Parameter Explanation' not in w or 'Data' not in w:raise ValueError('未发现Total参数字典/Data结构')
    fields={r[0]+1:{'field':r[3],'unit':r[4],'definition':r[6]} for r in w['Parameter Explanation'].iter_rows(min_row=2,values_only=True)
            if isinstance(r[0],int) and r[3]}
    if len({v['field'] for v in fields.values()})!=len(fields):raise ValueError('字段重名')
    ids=identity_values(w['Data']);sha=digest(path);rows=[]
    for row in w['Data'].iter_rows(min_row=7):
        if not any(c.value is not None for c in row):continue
        raw={fields[c.column]['field']:c.value for c in row if c.column in fields};ri=row[0].row
        for col in (1,2):
            if (ri,col) in ids:raw[fields[col]['field']]=ids[ri,col][0]
        rows.append({'sample_id':f'{sha}:Data:{ri}','source_row':ri,'original_values':raw})
    w.close();return rows,{v['field']:v for v in fields.values()}


def select_records(rows,scope,config):
    """存储Converted为目标；不补标签。算术检查不等于原文或系数已核实。"""
    records=[];selected=[]
    for r in rows:
        raw=r['original_values'];x=numeric(raw.get('CS_28d'));y=numeric(raw.get(config['target']));f=numeric(raw.get('Conversion Factor'))
        reasons=[]
        if raw.get('CastingEnv')!='Underwater':reasons.append('原表未标为Underwater')
        if not raw.get('PDF filename') or not raw.get('Research Group'):reasons.append('来源身份不明确')
        if x is None or x<=0:reasons.append('原始28天标签未知/非正值')
        if y is None or y<=0:reasons.append('Converted标签未知/非正值')
        if f not in config['accepted_provisional_factors']:reasons.append('不在本轮用户许可暂定系数规则内，未判定为坏数据')
        if x and y and f and abs(y-x*f)>config['target_consistency_tolerance_MPa']+1e-8:reasons.append('原始与换算强度算术矛盾超过容差')
        s=scope.get(r['sample_id'],{})
        state=s.get('scope_state','浇筑细节未知_仅原表Underwater')
        if state.startswith('非直接'):reasons.append('已知冲洗/渗透后重塑试件，单独隔离')
        direct=state.startswith('有直接')
        entry={'sample_id':r['sample_id'],'source_row':r['source_row'],'paper':raw.get('PDF filename'),
            'specimen_id':raw.get('Specimen ID'),'casting_environment':raw.get('CastingEnv'),
            'raw_CS_28d':x,'stored_Converted_CS_28d':y,'factor':f,'scope_state':state,
            'included_exploratory':not reasons,'included_primary':not reasons and direct,'reasons':reasons,
            'label_status':'探索性暂定换算，非正式已核实','formal_reliability_status':'待核查','density_assumption_status':'未知'}
        records.append(entry)
        if not reasons:selected.append({**r,'target':y,'direct_scope_evidence':direct,'label_status':entry['label_status']})
    return records,selected


def clusters(rows,rtol):
    """不看目标；配合比逐字段≤1%且缺失模式相同则保守连通，论文再与配比簇合组。"""
    n=len(rows);parent=list(range(n))
    def root(i):
        while parent[i]!=i:parent[i]=parent[parent[i]];i=parent[i]
        return i
    def union(a,b):
        a,b=root(a),root(b)
        if a!=b:parent[max(a,b)]=min(a,b)
    values=np.array([[numeric(r['original_values'].get(f)) if numeric(r['original_values'].get(f)) is not None else np.nan
                      for f in RECIPE_FIELDS] for r in rows],float)
    links=[]
    for i in range(n):
        if np.sum(values[i]>0)<2:continue # 不以共同未知、0占位形成虚假大簇
        for j in range(i):
            known=np.isfinite(values[i]);other=np.isfinite(values[j])
            if not np.array_equal(known,other) or np.sum(values[j]>0)<2:continue
            if np.all(np.isclose(values[i,known],values[j,known],rtol=rtol,atol=1e-8)):
                union(i,j);links.append({'row_a':rows[j]['source_row'],'row_b':rows[i]['source_row'],
                    'exact':bool(np.array_equal(values[i,known],values[j,known]))})
    def labels():
        roots=sorted({root(i) for i in range(n)});mapping={r:f'G{k:03}' for k,r in enumerate(roots)}
        return np.array([mapping[root(i)] for i in range(n)])
    recipes=labels();seen={}
    for i,r in enumerate(rows):
        for field in ['PDF filename','Research Group']:
            key=(field,r['original_values'][field])
            if key in seen:union(i,seen[key])
            else:seen[key]=i
    return recipes,labels(),links


def split_groups(groups,fraction,seed):
    t,v=next(GroupShuffleSplit(n_splits=1,test_size=fraction,random_state=seed).split(np.zeros(len(groups)),groups=groups))
    if set(groups[t])&set(groups[v]):raise ValueError('分组泄漏')
    return t,v


def model_pipeline(name,config):
    """固定参数，无搜索；模型缺依赖单独记为不可用，不阻断其他模型。"""
    p=config['parameters']
    if name=='rf':model=RandomForestRegressor(n_estimators=p['trees'],min_samples_leaf=3,max_features=1.,random_state=config['seed'],n_jobs=1)
    elif name=='xgb':
        from xgboost import XGBRegressor
        model=XGBRegressor(n_estimators=p['trees'],max_depth=p['max_depth'],learning_rate=p['learning_rate'],
            subsample=.85,colsample_bytree=.85,reg_lambda=5,objective='reg:squarederror',random_state=config['seed'],n_jobs=1)
    elif name=='catboost':
        from .nested_exploration import SklearnCatBoost
        model=SklearnCatBoost(iterations=p['trees'],depth=p['max_depth'],learning_rate=p['learning_rate'],l2_leaf_reg=5,
            loss_function='RMSE',random_seed=config['seed'],thread_count=1,verbose=False,allow_writing_files=False)
    elif name=='mean':model=DummyRegressor(strategy='mean')
    else:raise ValueError('未知模型')
    return Pipeline([('impute',SimpleImputer(strategy='median',add_indicator=True,keep_empty_features=True)),('model',model)])


def run(config_path,output):
    config_path,output=Path(config_path),Path(output);config=json.loads(config_path.read_text());output.mkdir(exist_ok=False)
    logger=(output/'training.log').open('x')
    def log(message):print(message,flush=True);logger.write(message+'\n');logger.flush()
    inv=inventory();write_json(output/'excel_inventory.json',inv)
    paths=[config_path,Path(__file__),Path(config['raw']),Path(config['scope_registry']),Path('src/concrete_audit/audit.py'),
        Path('src/concrete_audit/baselines.py'),Path('src/concrete_audit/parallel_exploration.py'),Path('src/concrete_audit/nested_exploration.py')]
    paths += [Path(r['path']) for r in inv['workbooks'] if r.get('kind')=='original_attachment']
    before={str(p.resolve()):digest(p) for p in paths}
    rows,dictionary=read_source(config['raw'])
    # 既有范围证据必须与本次从Excel重读的完整身份一致，不能仅按行号套用。
    scope={r['sample_id']:r for r in json.loads(Path(config['scope_registry']).read_text())['records']}
    for r in rows:
        if r['sample_id'] in scope:
            s=scope[r['sample_id']]
            if s['registry_pdf_filename']!=r['original_values']['PDF filename'] or s['specimen_id']!=r['original_values']['Specimen ID']:
                raise ValueError('范围证据身份不一致')
    records,selected=select_records(rows,scope,config)
    write_json(output/'selection_records.json',records);write_json(output/'field_dictionary.json',dictionary)
    statistics={'source_records':len(rows),'underwater_records':sum(r['original_values'].get('CastingEnv')=='Underwater' for r in rows),
        'positive_raw_28d_underwater':sum(r['original_values'].get('CastingEnv')=='Underwater' and (numeric(r['original_values'].get('CS_28d')) or 0)>0 for r in rows),
        'positive_converted_underwater':sum(r['original_values'].get('CastingEnv')=='Underwater' and (numeric(r['original_values'].get('Converted CS_28d')) or 0)>0 for r in rows),
        'primary_direct_evidence':sum(r['direct_scope_evidence'] for r in selected),'exploratory_total':len(selected),
        'unknown_scope_additional':sum(not r['direct_scope_evidence'] for r in selected),
        'missing_requested_workbooks':inv['missing_requested_workbooks']}
    write_json(output/'sample_statistics.json',statistics);log(json.dumps(statistics,ensure_ascii=False))
    candidates=config['features']+['SP','AWA','StF','SyF','w/c','w/b','Density']
    feature_info=[]
    for f in candidates:
        vals=[numeric(r['original_values'].get(f)) for r in selected]
        feature_info.append({'field':f,'selected':f in config['features'],'declared_unit':dictionary.get(f,{}).get('unit'),
            'positive_n':sum(v is not None and v>0 for v in vals),'zero_n':sum(v==0 for v in vals),
            'unknown_or_negative_n':sum(v is None or v<0 for v in vals),
            'reason':'输入零含义未知，训练折内插补' if f in config['features'] else
            '设计/实测属性未明确，不纳入' if f=='Density' else '单位/掺量基准未统一或公式来源待核，留待独立敏感性实验'})
    write_json(output/'features_and_missingness.json',feature_info)
    write_json(output/'preprocessing.json',{'features':config['features'],'target':config['target'],'rule':config['imputation'],
        'target_policy':'直接使用原表正值Converted；与原始CS×许可系数容差≤0.1MPa；不重算、不补标签，算术不授信',
        'density':'未纳入，未假定设计属性；源输入单位风险保留','zero':'所有所选材料≤0视未知；未声称为真实未掺',
        'random_split':'先按配比近重复簇随机划分，允许同论文不同配比跨集；不是普通逐行随机，避免近重复泄漏',
        'paper_split':'论文/研究组/跨论文近重复连通组隔离','tuning':'没有搜索或根据测试集修改参数'})
    write_json(output/'modeling_records.json',[{'sample_id':r['sample_id'],'source_row':r['source_row'],
        'paper':r['original_values']['PDF filename'],'target':r['target'],'direct_scope_evidence':r['direct_scope_evidence'],
        'label_status':r['label_status'],'features':dict(zip(config['features'],[None if np.isnan(v) else float(v) for v in feature_matrix([r],config['features'])[0]]))} for r in selected])
    results=[];predictions=[];splits=[];duplicate_reports=[];importances=[];unavailable=[]
    for cohort in ['primary_direct','exploratory_underwater']:
        data=[r for r in selected if r['direct_scope_evidence'] or cohort=='exploratory_underwater']
        X=feature_matrix(data,config['features']);y=np.array([r['target'] for r in data])
        recipes,papers,links=clusters(data,config['near_recipe_rtol'])
        duplicate_reports.append({'cohort':cohort,'n':len(data),'recipe_clusters':len(set(recipes)),
            'paper_connected_groups':len(set(papers)),'candidate_duplicate_pairs':links,
            'limitation':'1%逐字段与相同已知模式为操作性近重复规则，未证明识别所有跨文献重复；不按标签去重'})
        protocols=[]
        for mode,g in [('random_recipe_holdout',recipes),('paper_group_holdout',papers)]:
            t,v=split_groups(g,config['test_fraction'],config['seed']);protocols.append((mode,[(0,t,v)]))
        protocols.append(('paper_group_cv',[(k,t,v) for k,(t,v) in enumerate(grouped_folds(X,y,papers,min(config['group_cv_folds'],len(set(papers)))))]))
        for mode,folds in protocols:
            for k,t,v in folds:
                if set(recipes[t])&set(recipes[v]):raise ValueError('近重复配合比跨训练/验证')
                if mode.startswith('paper') and set(papers[t])&set(papers[v]):raise ValueError('论文跨训练/验证')
                splits.append({'cohort':cohort,'protocol':mode,'fold':k,'train_n':len(t),'test_n':len(v),
                    'train_ids':[data[i]['sample_id'] for i in t],'test_ids':[data[i]['sample_id'] for i in v],
                    'train_recipe_groups':sorted(set(recipes[t])),'test_recipe_groups':sorted(set(recipes[v])),
                    'train_paper_groups':sorted(set(papers[t])),'test_paper_groups':sorted(set(papers[v]))})
            for name in config['models']+['mean']:
                try:model_pipeline(name,config)
                except ImportError as exc:
                    unavailable.append({'model':name,'reason':str(exc)});continue
                pred=np.full(len(data),np.nan);train_scores=[];fold_scores=[]
                for k,t,v in folds:
                    model=model_pipeline(name,config);model.fit(X[t],y[t]);pv=model.predict(X[v]);pred[v]=pv
                    train_scores.append(regression_metrics(y[t],model.predict(X[t])));fold_scores.append({'fold':k,**regression_metrics(y[v],pv)})
                    for i,p in zip(v,pv):predictions.append({'cohort':cohort,'protocol':mode,'model':name,'fold':k,
                        'sample_id':data[i]['sample_id'],'source_row':data[i]['source_row'],'paper':data[i]['original_values']['PDF filename'],
                        'paper_group':str(papers[i]),'actual':float(y[i]),'predicted':float(p),'residual_actual_minus_predicted':float(y[i]-p)})
                    if mode.endswith('holdout'):
                        joblib.dump({'pipeline':model,'features':config['features'],'target':config['target'],'cohort':cohort,
                            'split':mode,'exploratory_only':True},output/f'{cohort}__{mode}__{name}.joblib')
                        if name!='mean':
                            perm=permutation_importance(model,X[v],y[v],scoring='neg_mean_absolute_error',n_repeats=3,random_state=config['seed'],n_jobs=1)
                            importances.append({'cohort':cohort,'protocol':mode,'model':name,
                                'MAE_increase_by_feature':dict(zip(config['features'],perm.importances_mean.tolist())),
                                'note':'本轮描述性测试集诊断；不能再用同一测试分数挑特征并称独立评价'})
                observed=np.isfinite(pred);gs=[regression_metrics(y[observed&(papers==g)],pred[observed&(papers==g)]) for g in sorted(set(papers[observed]))]
                result={'cohort':cohort,'protocol':mode,'model':name,'cohort_n':len(data),**regression_metrics(y[observed],pred[observed]),
                    'macro_source_MAE_MPa':float(np.mean([g['MAE_MPa'] for g in gs])),
                    'mean_training_MAE_MPa':float(np.mean([g['MAE_MPa'] for g in train_scores])),
                    'mean_training_R2':float(np.mean([g['R2'] for g in train_scores])),'folds':fold_scores}
                results.append(result);log(f'{cohort}/{mode}/{name} n={result["n"]} MAE={result["MAE_MPa"]:.4f} RMSE={result["RMSE_MPa"]:.4f} R2={result["R2"]:.4f}')
    for filename,value in [('results.json',results),('predictions.json',predictions),('split_manifest.json',splits),
        ('duplicate_analysis.json',duplicate_reports),('permutation_importance.json',importances),('unavailable_models.json',unavailable)]:write_json(output/filename,value)
    keys=['cohort','protocol','model','cohort_n','n','MAE_MPa','RMSE_MPa','R2','macro_source_MAE_MPa','mean_training_MAE_MPa','mean_training_R2']
    with (output/'metrics.csv').open('w',newline='',encoding='utf-8-sig') as f:
        writer=csv.DictWriter(f,fieldnames=keys,extrasaction='ignore');writer.writeheader();writer.writerows(results)
    snapshot=output/'code';snapshot.mkdir()
    for p in paths:
        if p.suffix=='.py' and p.is_file():shutil.copyfile(p,snapshot/p.name)
    shutil.copyfile(config_path,output/'config.json')
    after={p:digest(p) for p in before}
    if before!=after:raise RuntimeError('原始文件/输入发生变化')
    env={}
    for package in ['numpy','scikit-learn','openpyxl','xgboost-cpu','catboost','joblib']:
        try:env[package]=importlib.metadata.version(package)
        except importlib.metadata.PackageNotFoundError:env[package]='未安装'
    write_json(output/'manifest.json',{'experiment_id':config['experiment_id'],'created_utc':datetime.now(timezone.utc).isoformat(),
        'config':config,'input_hashes_before':before,'input_hashes_after':after,'environment':env,'statistics':statistics,
        'status':'exploratory_only','historical_data_use':config['historical_note'],'target_is_stored_converted':True,
        'no_target_imputation':True,'no_hyperparameter_search':True,'input_files_unchanged':True})
    logger.close();return results
