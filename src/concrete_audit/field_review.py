"""字段级本地辅助核查：比较经阅读整理的证据，不冒充外部AI自主判真。"""
from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
import math
from pathlib import Path
import re
import subprocess
from typing import Protocol
import openpyxl
from .reliability import number

STATUSES=('已核实','已核实但需转换','发现不一致','文献中未找到','证据不足','待人工确认')


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def normalized(text):
    """只规范空白和软连字符以定位摘录，不用近似匹配拼凑证据。"""
    return re.sub(r'\s+',' ',str(text).replace('\xad','')).strip()


def canonical_unit(unit):
    """仅合并等义单位写法；L与kg保持不同，不隐含密度。"""
    return {'kg/m³':'kg/m3','L/m³':'L/m3','l/m3':'L/m3','MPa':'MPa'}.get(unit,unit)


def match_value(a,b,tolerance=1e-9):
    na,nb=number(a),number(b)
    if na is not None and nb is not None:
        return abs(na-nb)<=tolerance
    return a==b


class ReviewBackend(Protocol):
    """可替换后端；新增外部模型时仍需经过证据校验及人工复核。"""
    def review(self,current,unit,claim): ...


class LocalEvidenceBackend:
    """对经人工式阅读编录的字段证据做确定性检查；不是大模型调用。"""
    def review(self,current,unit,claim):
        if claim is None:
            return {'status':'证据不足','difference':'本轮尚未编录该字段证据；不等于论文未报告。',
                    'suggestion':'继续定位对应配方、环境、龄期与字段。','value_category':'C','computed_value':None}
        mode=claim['mode']
        result={'computed_value':None,'value_category':'C','difference':claim.get('note',''),
                'suggestion':claim.get('action','保留原值，待人工复核。')}
        if mode=='direct':
            if canonical_unit(unit)!=canonical_unit(claim['unit']):
                result.update(status='已核实但需转换',value_category='C',
                              difference='原文值已定位，但原文与Excel单位不同；不得直接比较数字或假设密度。 '+result['difference'])
            else:
                same=match_value(current,claim['comparison_value'])
                result.update(status='已核实' if same else '发现不一致',value_category='A' if same else 'C')
                if not same:result['difference']='Excel值与已定位原文值不一致。 '+result['difference']
        elif mode=='calculated':
            if claim.get('formula_id')=='percent_mass':
                pct=number(claim['parameters']['percent']);mass=number(claim['parameters']['mass_kg_m3'])
                if pct is None or mass is None or pct<0 or mass<=0:
                    raise ValueError('计算输入无效')
                result['computed_value']=pct/100*mass
                result['value_category']='B（候选计算，未审批）'
                result['difference']+=f' 按原文参数计算={result["computed_value"]:.8g} kg/m3；不回写原表。'
            elif claim.get('formula_id')=='sum_components':
                values=[number(x) for x in claim['parameters']['components_kg_m3']]
                if not values or any(x is None or x<0 for x in values):raise ValueError('分项质量无效')
                result['computed_value']=sum(values)
                result['value_category']='B（候选计算，未审批）'
                result['difference']+=f' 原文分项合计={result["computed_value"]:.8g} kg/m3；不回写原表。'
            result['status']='已核实但需转换'
        elif mode in ('conflict','ambiguous','insufficient','not_found'):
            result['status']={'conflict':'发现不一致','ambiguous':'待人工确认',
                              'insufficient':'证据不足','not_found':'文献中未找到'}[mode]
            if mode=='not_found' and not claim.get('search_scope'):
                raise ValueError('未找到结论必须声明检索范围')
        else:
            raise ValueError('未知核查模式')
        return result


def validate_evidence(claim,pages,expected_pdf_sha256=None):
    """页码、摘录均绑定实际提取文本；匹配不等于证明字段解释正确。"""
    if not claim.get('evidence'):
        raise ValueError('已编录结论必须附实际证据')
    for ref in claim['evidence']:
        page=ref['page'];quote=ref['quote']
        if not isinstance(page,int) or page<1 or page>len(pages) or not quote.strip():
            raise ValueError('证据页码或摘录无效')
        if normalized(quote) not in normalized(pages[page-1]):
            raise ValueError('摘录不在指定PDF页的实际提取文本中')
        if ref.get('image_path'):
            if not expected_pdf_sha256 or ref.get('source_pdf_sha256')!=expected_pdf_sha256:
                raise ValueError('图像证据的PDF来源指纹不匹配')
            if digest(ref['image_path'])!=ref.get('image_sha256'):
                raise ValueError('图像证据内容指纹不匹配')


def run(config_path,output,backend=None):
    """全新目录运行，保留每项证据与比较；不改变原始值或可靠状态。"""
    config_path,output=Path(config_path),Path(output)
    config=json.loads(config_path.read_text())
    protected=[config_path,Path(config['raw']),Path(config['annotated']),Path(config['registry']),
               *[Path(x['path']) for x in config['papers'].values()]]
    before={str(p):digest(p) for p in protected}
    for paper in config['papers'].values():
        if digest(paper['path'])!=paper['sha256']:raise ValueError('PDF版本指纹不匹配')
    if digest(config['raw'])!=config['raw_sha256']:raise ValueError('原始Excel版本不匹配')
    output.mkdir(parents=True,exist_ok=False)
    backend=backend or LocalEvidenceBackend()
    pdf_pages={}
    for key,paper in config['papers'].items():
        target=output/(key+'.txt')
        subprocess.run(['pdftotext','-layout',paper['path'],str(target)],check=True,capture_output=True)
        pages=target.read_text().split('\f')
        if not pages[-1].strip():pages.pop()
        pdf_pages[key]=pages
    raw=openpyxl.load_workbook(config['raw'],data_only=False)
    annotated=openpyxl.load_workbook(config['annotated'],data_only=False)
    params={r[3]:{'column':r[0]+1,'unit':r[4],'definition':r[6]}
            for r in raw['Parameter Explanation'].iter_rows(min_row=2,values_only=True)
            if isinstance(r[0],int) and r[3]}
    registry={r['source_row']:r for r in json.loads(Path(config['registry']).read_text())}
    checks=[];samples=[];seen=set()
    for sample in config['samples']:
        row=sample['source_row'];key=sample['paper'];spec=raw['Data'].cell(row,params['Specimen ID']['column']).value
        if spec!=sample['specimen_id']:raise ValueError('样本身份不匹配')
        if row in seen:raise ValueError('重复样本行号')
        seen.add(row)
        # 校验来源映射不能只依赖行号或作者姓名。
        if registry[row]['pdf_filename']!=sample['registry_pdf_filename']:
            raise ValueError('原档案论文身份不匹配')
        fields=sample['fields']
        if len(fields)!=len(set(fields)):raise ValueError('重复样本字段')
        for field in fields:
            claim=sample['claims'].get(field)
            if claim:validate_evidence(claim,pdf_pages[key],config['papers'][key]['sha256'])
            if field not in params:raise ValueError(f'字段不存在：{field}')
            info=params[field];cell=raw['Data'].cell(row,info['column'])
            current=cell.value;aux=annotated['Data'].cell(row,info['column']).value
            result=backend.review(current,info['unit'],claim)
            checks.append({'check_id':f'{config["raw_sha256"]}:Data:{row}:{field}',
                'sample_id':f'{config["raw_sha256"]}:Data:{row}',
                'source_row':row,'specimen_id':spec,'field':field,'excel_cell':cell.coordinate,
                'excel_value':current,'annotated_value':aux,'excel_unit':info['unit'],
                'field_definition':info['definition'],'paper':key,
                'pdf_path':config['papers'][key]['path'],'pdf_sha256':config['papers'][key]['sha256'],
                'paper_value':claim.get('paper_value') if claim else None,
                'paper_unit':claim.get('unit') if claim else None,
                'evidence':claim['evidence'] if claim else [],
                'search_scope':claim.get('search_scope') if claim else '本轮尚未编录',
                'formula':claim.get('formula') if claim else None,
                'parameters':claim.get('parameters') if claim else None,
                'extraction_method':claim.get('method','文本/表格阅读编录') if claim else '尚未核实',
                'comparison_backend':'local_evidence_annotations_v1',
                'annotation_author':'本会话助手阅读编录，未经用户人工签字',
                'status_scope':'字段级辅助结论，尚未经人工最终确认',
                'requires_human_confirmation':True,'human_confirmation':'待确认',**result})
        samples.append({'source_row':row,'specimen_id':spec,'paper':key,
                        'candidate_in_527':registry[row]['candidate_link'],
                        'original_registry_status':registry[row]['status'],
                        'status_changed':False,'new_trusted_qualification':False})
    counts={s:sum(r['status']==s for r in checks) for s in STATUSES}
    summary={'samples':len(samples),'papers':len(pdf_pages),'field_checks':len(checks),
             'candidate_samples':sum(s['candidate_in_527'] for s in samples),'status_counts':counts,
             'fields_with_located_evidence':sum(bool(r['evidence']) for r in checks),
             'fields_not_yet_annotated':sum(not r['evidence'] for r in checks),
             'new_trusted_samples':0,'external_llm_called':False,'model_training':False,
             'all_require_human_confirmation':True,'inputs_unchanged':all(digest(Path(p))==v for p,v in before.items())}
    if not summary['inputs_unchanged']:raise RuntimeError('输入发生变化')
    human=[r for r in checks if r['status']!='已核实']
    for name,obj in [('field_checks.json',checks),('sample_index.json',samples),('human_review.json',human),('summary.json',summary),
                     ('manifest.json',{'inputs_sha256':before,'created_at_utc':datetime.now(timezone.utc).isoformat(),
                                       'source_code_sha256':digest(Path(__file__))})]:
        (output/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    workbook=openpyxl.Workbook();workbook.remove(workbook.active)
    for title,data,keys in [('字段核查',checks,['source_row','specimen_id','field','excel_value','annotated_value','excel_unit','paper_value','paper_unit','status','difference','suggestion','computed_value','formula','parameters','paper','evidence','requires_human_confirmation','human_confirmation','status_scope','pdf_sha256']),
                            ('待人工核查',human,['source_row','field','status','difference','suggestion','paper','evidence']),
                            ('样本索引',samples,list(samples[0]))]:
        sheet=workbook.create_sheet(title);sheet.append(keys)
        for item in data:
            sheet.append([json.dumps(item.get(k),ensure_ascii=False) if isinstance(item.get(k),(list,dict)) else item.get(k) for k in keys])
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for col in sheet.columns:sheet.column_dimensions[col[0].column_letter].width=24
    workbook.save(output/'小批字段核查.xlsx')
    workbook.close();raw.close();annotated.close()
    return summary
