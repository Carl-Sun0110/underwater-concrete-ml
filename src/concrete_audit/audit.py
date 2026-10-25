"""只读解析工作簿、保留批注和格式，输出独立审计档案。"""
import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
import importlib.metadata
import json
from pathlib import Path
import re
import sys

import openpyxl
from openpyxl.utils.cell import range_boundaries
from .reliability import UNKNOWN, PENDING, EXCLUDED, conversion_check, apply_review, number


def sha256(path):
    """计算内容指纹，用于版本绑定及只读验证。"""
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def dump(path, value):
    """仅在新运行目录中创建文件，拒绝覆盖。"""
    with Path(path).open('x', encoding='utf-8') as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2, default=str, allow_nan=False)


def table(sheet):
    """读取标准首行表头；拒绝重复或空字段名。"""
    rows = list(sheet.iter_rows(values_only=True))
    header = rows[0]
    if any(x is None for x in header) or len(set(header)) != len(header):
        raise ValueError(f'{sheet.title} 表头不唯一或为空')
    return [dict(zip(header, row)) for row in rows[1:] if any(x is not None for x in row)]


def identity_values(sheet):
    """仅展开A/B列已存在的合并单元格；普通空白不擅自向下补齐。"""
    values = {}
    for area in sheet.merged_cells.ranges:
        if area.min_col == area.max_col and area.min_col in (1, 2):
            anchor = sheet.cell(area.min_row, area.min_col)
            for row in range(max(7, area.min_row), area.max_row+1):
                values[(row, area.min_col)] = (anchor.value, anchor.coordinate)
    return values


def comment_rows(text, anchor_row):
    """关联批注明确写出的A1:B2/A1～B2范围；不解释自然语言或核查结论。"""
    rows = {anchor_row}
    for match in re.finditer(r'\b([A-Z]{1,3}\d+)\s*[:：～~—–]\s*([A-Z]{1,3}\d+)\b', text):
        _, low, _, high = range_boundaries(match[1]+':'+match[2])
        if low <= high and high <= 100000:
            rows.update(range(low, high+1))
    return rows


def same(a, b):
    """允许整数/小数的等值表示，不容忍实际数值变化。"""
    na, nb = number(a), number(b)
    return (na == nb) if na is not None and nb is not None else a == b


def run(raw_path, screening_path, output, evidence_paths=(), decisions_path=None):
    """执行审计；每次使用全新输出目录，输入工作簿从不保存回写。"""
    raw_path, screening_path, output = map(Path, (raw_path, screening_path, output))
    inputs = [raw_path, screening_path, *map(Path, evidence_paths)]
    if decisions_path:
        inputs.append(Path(decisions_path))
    if len({p.resolve() for p in inputs}) != len(inputs):
        raise ValueError('输入文件重复')
    before = {str(p.resolve()): sha256(p) for p in inputs}
    output.mkdir(parents=True, exist_ok=False)
    raw = openpyxl.load_workbook(raw_path, data_only=False)
    screening = openpyxl.load_workbook(screening_path, read_only=True, data_only=False)
    for name in ('Data', 'Parameter Explanation'):
        if name not in raw:
            raise ValueError(f'原始工作簿缺少{name}')
    for name in ('初筛候选527', '待核查清单', '特征完整性', '筛选规则', '样本统计'):
        if name not in screening:
            raise ValueError(f'筛选工作簿缺少{name}')
    sheet = raw['Data']
    # 参数字典是静态表；Data第4行是公式表头，不能要求Excel计算器参与审计。
    params = {}
    for row in raw['Parameter Explanation'].iter_rows(min_row=2, values_only=True):
        if isinstance(row[0], int) and row[3]:
            params[row[0]+1] = {'field': row[3], 'unit': row[4], 'definition': row[6]}
    if len({v['field'] for v in params.values()}) != len(params):
        raise ValueError('参数字典字段重名')
    required = {'PDF filename','Research Group','Specimen ID','CastingEnv','CS_28d',
                'Conversion Factor','Converted CS_28d','ConvertFlag','CS_Spec','Density'}
    if not required <= {v['field'] for k,v in params.items() if k <= sheet.max_column}:
        raise ValueError('原始工作簿缺少必要字段')
    candidates = table(screening['初筛候选527'])
    reviews = table(screening['待核查清单'])
    for label, rows in [('候选', candidates), ('核查', reviews)]:
        keys = [r['原表行号'] for r in rows]
        if len(keys) != len(set(keys)):
            raise ValueError(f'{label}原表行号重复，禁止歧义连接')
    candidate_map = {r['原表行号']: r for r in candidates}
    review_map = {r['原表行号']: r for r in reviews}
    comments, styles, style_catalog = [], [], {}
    attached = {}
    for ws in raw:
        for row in ws:
            for cell in row:
                if cell.comment:
                    entry = {'id': f'{ws.title}!{cell.coordinate}', 'sheet': ws.title,
                             'cell': cell.coordinate, 'text': cell.comment.text,
                             'interpretation': '附件中的核查陈述，未独立验证'}
                    comments.append(entry)
                    if ws.title == 'Data':
                        for r in comment_rows(cell.comment.text, cell.row):
                            attached.setdefault(r, []).append(entry['id'])
    merged = identity_values(sheet)
    registry, diffs, links = [], [], []
    observed = set()
    raw_hash = before[str(raw_path.resolve())]
    for row in sheet.iter_rows(min_row=7):
        if not any(c.value is not None for c in row):
            continue
        rownum = row[0].row
        observed.add(rownum)
        values = {params[c.column]['field']: c.value for c in row if c.column in params}
        identities = {}
        for col in (1, 2):
            if (rownum, col) in merged:
                val, anchor = merged[(rownum, col)]
                values[params[col]['field']] = val
                identities[params[col]['field']] = f'合并单元格锚点{anchor}'
            else:
                identities[params[col]['field']] = f'原单元格{sheet.cell(rownum,col).coordinate}'
        candidate = candidate_map.get(rownum)
        review = review_map.get(rownum)
        matched = {}
        for kind, external in [('candidate',candidate),('review',review)]:
            if external is None:
                matched[kind] = False
                continue
            ok = all(same(values.get(k),external.get(k)) for k in ('PDF filename','Research Group','Specimen ID'))
            matched[kind] = ok
            links.append({'source_row':rownum,'table':kind,'identity_match':ok})
            for key in sorted(external.keys() & values.keys()):
                if not same(values[key],external[key]):
                    diffs.append({'source_row':rownum,'table':kind,'field':key,
                                  'raw_value':values[key],'screening_value':external[key],
                                  'identity_match':ok,'action':'仅报告，不覆盖任一版本'})
        reasons = ['来源论文证据尚未独立核实', '密度假设状态未知', '输入单位及零值语义尚未全部核实']
        check = conversion_check(values)
        reasons.extend(check['reasons'])
        for kind, external in [('candidate',candidate),('review',review)]:
            if external and not matched[kind]:
                reasons.append(f'{kind}身份不匹配，未自动关联其核查结论')
        if matched['candidate'] and candidate.get('待核查问题'):
            reasons.append('旧标记：'+str(candidate['待核查问题']))
        if matched['review']:
            reasons.append('筛选工作簿：'+str(review.get('本轮待核查原因')))
        if attached.get(rownum):
            reasons.append('存在批注陈述，须逐字段核实并确认范围')
        status = EXCLUDED if values.get('CastingEnv') == 'Air' else PENDING
        if status == EXCLUDED:
            reasons.append('明确为Air，按本阶段仅水下浇筑范围排除；不代表数据本身错误')
        elif values.get('CastingEnv') != 'Underwater':
            reasons.append('浇筑环境未知或未标准化')
        record = {'record_id': f'{raw_hash}:Data:{rownum}', 'source_sha256':raw_hash,
                  'source_sheet':'Data', 'source_row':rownum,
                  'pdf_filename':values.get('PDF filename'), 'research_group':values.get('Research Group'),
                  'specimen_id':values.get('Specimen ID'), 'identity_provenance':identities,
                  'casting_env':values.get('CastingEnv'), 'data_source':values.get('DataSource'),
                  'candidate_link':matched['candidate'], 'review_link':matched['review'],
                  'historical_status':review.get('旧筛选状态') if matched['review'] else
                                      candidate.get('旧筛选状态') if matched['candidate'] else UNKNOWN,
                  'status':status, 'strength_basis_status':UNKNOWN, 'density_assumption_status':UNKNOWN,
                  'source_evidence_status':UNKNOWN, 'feature_units_status':UNKNOWN,
                  'zero_semantics_status':UNKNOWN,'conversion':check,
                  'strength_fields':{k:values.get(k) for k in ('CS_28d','Conversion Factor','Converted CS_28d','CS_Spec','CS_Standard')},
                  'density_raw':values.get('Density'), 'pending_reasons':reasons,
                  'comment_ids':attached.get(rownum,[]),
                  'processing_history':[{'action':'只读导入及机械校验','state':status,
                                         'rule':'Air仅按研究范围排除；其他记录默认待核查，旧排除仅保留为历史状态'}]}
        registry.append(record)
        for cell in row:
            if cell.has_style and cell.value is not None:
                styles.append({'cell':cell.coordinate,'style_id':cell.style_id})
                if cell.style_id not in style_catalog:
                    style_catalog[cell.style_id] = {
                        'fill_type':cell.fill.patternType,'fill_color':str(cell.fill.fgColor),
                        'font_color':str(cell.font.color),'bold':cell.font.bold,
                        'interpretation':'原始格式线索，不自动判定可靠性'}
    orphan = sorted((set(candidate_map)|set(review_map))-observed)
    events = []
    if decisions_path:
        decisions = json.loads(Path(decisions_path).read_text(encoding='utf-8'))
        by_id = {r['record_id']: i for i,r in enumerate(registry)}
        seen = set()
        evidence_hashes = {sha256(p) for p in evidence_paths}
        for decision in decisions:
            key = decision['record_id']
            if key not in by_id or key in seen:
                raise ValueError('人工决策指向未知或重复记录')
            seen.add(key)
            i = by_id[key]
            updated = apply_review(registry[i],decision,evidence_hashes)
            event = {'record_id':key,'from':registry[i]['status'],'to':updated['status'],'decision':decision}
            updated['processing_history'].append(event)
            registry[i] = updated
            events.append(event)
    completeness = []
    for expected in table(screening['特征完整性']):
        field = expected['字段']; nums = [number(r.get(field)) for r in candidates]
        actual = [len(nums),sum(x is not None for x in nums),sum(x==0 for x in nums),
                  sum(x is not None and x>0 for x in nums),sum(x is None for x in nums)]
        reported = [expected[k] for k in ('候选总数','数值记录数','数值为0记录数','正值记录数','占位/非数值记录数')]
        completeness.append({'field':field,'actual':actual,'reported':reported,'matches':actual==reported})
    sheets = [{'file':str(p),'sheets':[{'name':s.title,'rows':s.max_row,'columns':s.max_column} for s in w]}
              for p,w in [(raw_path,raw),(screening_path,screening)]]
    evidence_inventory = [{'pdf_filename': name,
                           'record_count':sum(r['pdf_filename']==name for r in registry),
                           'provided_file_matches':[str(p) for p in evidence_paths if Path(p).name==name],
                           'independently_verified':False}
                          for name in sorted({r['pdf_filename'] for r in registry if r['pdf_filename']})]
    raw_underwater = sum(r['casting_env']=='Underwater' for r in registry)
    sample_rows = list(screening['样本统计'].values)
    prior_underwater = next((r[1] for r in sample_rows if r[0]=='CastingEnv=Underwater'),None)
    known_missing = ['水下混凝土_人工核查版.xlsx（未提供同名版本；不认定批注版与其等同）',
                     '水下混凝土_筛选结果_20261012.xlsx（目前只有其衍生汇总）',
                     '来源论文PDF及页码/表格/图像证据', '完整强度换算依据及适用条件',
                     '逐条密度假设、单位和零值语义的人工核查决策档案']
    after = {str(p.resolve()):sha256(p) for p in inputs}
    if before != after:
        raise RuntimeError('输入文件指纹发生变化，停止生成成功报告')
    summary = {'raw_records':len(registry),'raw_unique_pdf_sources':len(evidence_inventory),
               'historical_underwater_reported':prior_underwater,
               'underwater_count_matches_historical':raw_underwater==prior_underwater,'casting_env_counts':dict(Counter(r['casting_env'] for r in registry)),
               'candidate_records_in_screening':len(candidates),
               'candidate_status_counts':dict(Counter(r['status'] for r in registry if r['candidate_link'])),
               'candidate_field_differences':sum(r['table']=='candidate' for r in diffs),
               'matched_candidate_records':sum(r['candidate_link'] for r in registry),
               'review_records_in_screening':len(reviews),
               'status_counts':{s:sum(r['status']==s for r in registry) for s in ('待核查','已核实可用','已排除')},
               'comments':len(comments),'cross_version_field_differences':len(diffs),
               'identity_mismatch_links':sum(not r['identity_match'] for r in links),
               'orphan_source_rows':orphan,'completeness_all_match':all(r['matches'] for r in completeness),
               'inputs_unchanged':before==after,'missing_files_or_evidence':known_missing}
    for name,obj in [('registry.json',registry),('comments.json',comments),('cell_styles.json',{'catalog':style_catalog,'cells':styles}),
                     ('version_differences.json',diffs),('evidence_inventory.json',evidence_inventory),('links.json',links),('completeness.json',completeness),
                     ('workbook_structure.json',sheets),('parameter_dictionary.json',params),('summary.json',summary),
                     ('manual_review_events.json',events),
                     ('screening_rules_snapshot.json',list(screening['筛选规则'].values)),
                     ('sample_statistics_snapshot.json',list(screening['样本统计'].values))]:
        dump(output/name,obj)
    dump(output/'manifest.json',{'inputs_sha256':before,'inputs_sha256_after':after,
         'created_at_utc':datetime.now(timezone.utc).isoformat(),'python':sys.version,
         'openpyxl':importlib.metadata.version('openpyxl'),'arguments':{
          'raw':str(raw_path),'screening':str(screening_path),'evidence':list(map(str,evidence_paths)),
          'decisions':str(decisions_path) if decisions_path else None},
         'source_code_sha256':{p.name:sha256(p) for p in Path(__file__).parent.glob('*.py')}})
    text = ['# 模块1与模块2审计报告','', '本报告仅审计及管理可靠状态；未预处理、未训练模型。', '',
            '## 运行结果','', '```json',json.dumps(summary,ensure_ascii=False,indent=2),'```','',
            '## 版本差异与待补文件','',
            f'批注版实际计数为水下{raw_underwater}条；历史统计为{prior_underwater}条。缺少历史原表，不能定位或解释全部差额。',
            f'候选身份成功关联{summary["matched_candidate_records"]}条；候选共享字段差异{summary["candidate_field_differences"]}项。',
            f'跨版本差异按字段计数：{dict(Counter(x["field"] for x in diffs))}；逐项见version_differences.json，未覆盖任一版本。',
            f'{len(evidence_inventory)}个主表PDF文件名仅为来源引用，文件提供情况见evidence_inventory.json；程序未独立核实论文内容。',
            *['- '+item for item in known_missing], '',
            '## 解释与限制','',
            '注意：批注版当前水下/空气计数与历史汇总可能不同，请查看summary.json，不能据此覆盖历史版本。',
            '已核实可用=0表示本次没有提供通过证据门槛的整条人工核查决策，不表示没有潜在可用数据。',
            '已排除仅表示原表明确为Air而不在本研究范围；历史排除状态保留，未自动升级为当前最终结论。',
            '批注是附件中的核查陈述，并非本程序亲自查阅论文的结果；即使写有“已核对”也不自动授予整条可信资格。',
            '批注只关联锚点及明确A1:B2形式的行范围；其余自然语言范围及字段适用性需人工确认。',
            '仅展开原有身份合并单元格，其他空白保持未知；跨版本先核对三项身份，再保留差异，不覆盖数据。',
            '机械换算一致不代表换算系数有可靠依据；零值未填补，Density未据正负判定来源。',
            '缺少的同名原始文件不能由批注版或衍生工作簿自动替代；尚不能完整重现历史筛选流程。','',
            '## 输出索引','',
            '- registry.json：逐条来源、状态、未知项、机械结果、待核查原因和处理历史。',
            '- comments.json / cell_styles.json：批注全文、位置及原始格式线索。',
            '- version_differences.json / links.json：跨版本字段差异与身份关联检查。',
            '- completeness.json：特征完整性逐项复算。',
            '- workbook_structure.json / parameter_dictionary.json：结构与单位字典。',
            '- manifest.json：输入路径、内容哈希、代码指纹及运行环境。','',
            '下一步：先审核版本差异和批注，补充论文证据与人工决策；确认后再进入预处理。']
    with (output/'report.md').open('x',encoding='utf-8') as stream:
        stream.write('\n'.join(text)+'\n')
    raw.close(); screening.close()
    return summary


def main():
    """显式指定输入和新输出目录，避免通过模糊文件名意外选择版本。"""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--raw',required=True,type=Path)
    parser.add_argument('--screening',required=True,type=Path)
    parser.add_argument('--output',required=True,type=Path)
    parser.add_argument('--evidence',nargs='*',default=[],type=Path)
    parser.add_argument('--decisions',type=Path)
    args = parser.parse_args()
    print(json.dumps(run(args.raw,args.screening,args.output,args.evidence,args.decisions),ensure_ascii=False,indent=2))


if __name__ == '__main__':
    main()
