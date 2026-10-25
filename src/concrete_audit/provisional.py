"""用户暂定系数的探索性标签；与原始值、文献核查和正式可信资格分离。"""
from collections import Counter
from datetime import datetime,timezone
from decimal import Decimal, InvalidOperation
import hashlib
import json
from pathlib import Path
import sys
import openpyxl
from .audit import identity_values, table


TARGET='Provisional_Converted_CS_28d'


def finite_decimal(value):
    """按Excel标量的十进制文本计算；不读公式、不把布尔/占位当数值。"""
    if isinstance(value,bool):return None
    try:
        n=Decimal(str(value))
        return n if n.is_finite() else None
    except (InvalidOperation,ValueError,TypeError):return None


def provisional_strength(strength,factor,allowed):
    """仅应用已有且在白名单的系数；缺标签不填补，缺系数不按规格推断。"""
    x,k=finite_decimal(strength),finite_decimal(factor)
    accepted={finite_decimal(v) for v in allowed}
    if None in accepted or any(v<=0 for v in accepted):raise ValueError('无效系数白名单')
    reasons=[]
    if x is None or x<=0:reasons.append('原始28d强度缺失、无效或非正')
    if k is None:reasons.append('系数缺失或非数值')
    elif k not in accepted:reasons.append('原有系数不在本次授权白名单')
    if reasons:return None,reasons
    return x*k,[]


def hash_file(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def eligible_reasons(row,candidate_match,label_reasons):
    """决定探索性标签子集，不判断原文是否可信；保留所有未纳入理由。"""
    reasons=list(label_reasons)
    if row.get('CastingEnv')!='Underwater':reasons.append('表内未标为Underwater')
    if not candidate_match:reasons.append('未与原527候选身份一致关联')
    if finite_decimal(row.get('ConvertFlag'))!=Decimal(0):reasons.append('ConvertFlag未明确为0')
    if row.get('CS_Spec') in (None,'','-','/',0):reasons.append('表内试件规格缺失或占位')
    return reasons


def run(inputs_path,policy_path,output):
    """只读导入Total；每次新目录输出所有记录及探索子集，不更新可信主档。"""
    inputs_path,policy_path,output=map(Path,(inputs_path,policy_path,output))
    inputs=json.loads(inputs_path.read_text());policy=json.loads(policy_path.read_text())
    if policy.get('promote_reliability_status') is not False or policy.get('fill_missing_strength') is not False:
        raise ValueError('探索规则必须禁止授信及目标填补')
    paths=[inputs_path,policy_path,*[Path(inputs[k]) for k in ('raw','annotated','screening','registry')]]
    before={str(p.resolve()):hash_file(p) for p in paths}
    registry={x['source_row']:x for x in json.loads(Path(inputs['registry']).read_text())}
    raw=openpyxl.load_workbook(inputs['raw'],data_only=False)
    annotated=openpyxl.load_workbook(inputs['annotated'],data_only=False,read_only=True)
    screening=openpyxl.load_workbook(inputs['screening'],data_only=False,read_only=True)
    params={r[0]+1:r[3] for r in raw['Parameter Explanation'].iter_rows(min_row=2,values_only=True) if isinstance(r[0],int) and r[3]}
    if len(set(params.values()))!=len(params):raise ValueError('原始字段名不唯一')
    candidates=table(screening['初筛候选527'])
    by_candidate={r['原表行号']:r for r in candidates}
    if len(by_candidate)!=len(candidates):raise ValueError('候选行号重复')
    required={'CS_28d','Conversion Factor','CastingEnv','CS_Spec','ConvertFlag','PDF filename','Research Group','Specimen ID'}
    if not required<=set(params.values()):raise ValueError('缺少必要字段')
    # 将批注版仅作为独立参照，不替换原版。
    annotated_rows={i:dict(zip((params[k] for k in sorted(params) if k<=annotated['Data'].max_column),r))
                    for i,r in enumerate(annotated['Data'].iter_rows(min_row=7,values_only=True),7)}
    identities=identity_values(raw['Data'])
    data=[]
    input_hash=before[str(Path(inputs['raw']).resolve())]
    candidate_identities=set()
    for row in raw['Data'].iter_rows(min_row=7):
        if not any(c.value is not None for c in row):continue
        ri=row[0].row;values={params[c.column]:c.value for c in row if c.column in params}
        for col in (1,2):
            if (ri,col) in identities:values[params[col]]=identities[ri,col][0]
        reg=registry.get(ri)
        if not reg:raise ValueError(f'缺少原可靠状态档案行{ri}')
        if (reg['pdf_filename'],reg['specimen_id'])!=(values['PDF filename'],values['Specimen ID']):
            raise ValueError(f'原档案身份冲突{ri}')
        cand=by_candidate.get(ri)
        matched=bool(cand and all(cand[k]==values[k] for k in ('PDF filename','Research Group','Specimen ID')))
        if cand and not matched:raise ValueError(f'候选身份不匹配{ri}')
        if matched:candidate_identities.add(ri)
        result,reasons=provisional_strength(values['CS_28d'],values['Conversion Factor'],policy['accepted_existing_row_factors'])
        excluded=eligible_reasons(values,matched,reasons)
        old=finite_decimal(values.get('Converted CS_28d'))
        delta=abs(result-old) if result is not None and old is not None else None
        entry={
            'sample_id':f'{input_hash}:Data:{ri}','source_row':ri,
            'original_values':values,
            'annotated_strength':{k:annotated_rows[ri].get(k) for k in ('CS_28d','Conversion Factor','Converted CS_28d','CS_Spec')},
            TARGET:float(result) if result is not None else None,
            'provisional_target_decimal':str(result) if result is not None else None,
            'difference_from_original_converted':float(delta) if delta is not None else None,
            'difference_exceeds_0_1_mpa':delta>Decimal('.1') if delta is not None else None,
            'candidate_in_original_527':matched,
            'exploratory_label_subset':not excluded,
            'not_in_exploratory_subset_reasons':excluded,
            'label_calculation_status':'暂定计算，非文献已核实' if result is not None else '未按本规则计算',
            'scope_basis':'依据Excel成型环境字段，非本轮原文逐条确认',
            'prior_reliability_status':reg['status'],
            'prior_pending_reasons':reg['pending_reasons'],
            'density_assumption_status':reg['density_assumption_status'],
            'strength_basis_status':policy['strength_basis_status'],
            'policy_id':policy['policy_id'],
            'training_authorized':False,'trusted_qualification_granted':False,
        }
        data.append(entry)
    if len(candidate_identities)!=len(candidates):raise ValueError('部分候选行在原表缺失')
    output.mkdir(parents=True,exist_ok=False)
    eligible=[r for r in data if r['exploratory_label_subset']]
    candidates_out=[r for r in data if r['candidate_in_original_527']]
    summary={
        'total_records':len(data),'underwater_marked_in_excel':sum(r['original_values']['CastingEnv']=='Underwater' for r in data),
        'original_candidate_count':len(candidates_out),'calculated_labels_all_environments':sum(r[TARGET] is not None for r in data),
        'calculated_labels_underwater':sum(r[TARGET] is not None and r['original_values']['CastingEnv']=='Underwater' for r in data),
        'exploratory_label_subset_count':len(eligible),'original_candidates_not_in_subset':len(candidates_out)-len(eligible),
        'candidate_factors':dict(Counter(str(r['original_values']['Conversion Factor']) for r in candidates_out)),
        'eligible_factors':dict(Counter(str(r['original_values']['Conversion Factor']) for r in eligible)),
        'eligible_research_groups':len({r['original_values']['Research Group'] for r in eligible}),
        'eligible_pdf_sources':len({r['original_values']['PDF filename'] for r in eligible}),
        'candidate_noninclusion_reasons_nonexclusive':dict(Counter(x for r in candidates_out if not r['exploratory_label_subset'] for x in r['not_in_exploratory_subset_reasons'])),
        'new_trusted_samples':0,'original_inputs_unchanged':all(hash_file(Path(k))==v for k,v in before.items()),
        'model_training':False,'feature_preprocessing':False,
    }
    if not summary['original_inputs_unchanged']:raise RuntimeError('输入文件变化')
    for name,obj in [('all_records.json',data),('exploratory_labels.json',eligible),('summary.json',summary),
                     ('policy_snapshot.json',policy),('manifest.json',{'inputs_sha256':before,'created_at_utc':datetime.now(timezone.utc).isoformat(),
                                            'source_code_sha256':hash_file(Path(__file__)),'python':sys.version})]:
        (output/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2,default=str,allow_nan=False),encoding='utf-8')
    workbook=openpyxl.Workbook();workbook.remove(workbook.active)
    fields=['source_row','PDF filename','Research Group','Specimen ID','CastingEnv','CS_28d','CS_Spec','Conversion Factor','Converted CS_28d',TARGET,
            'provisional_target_decimal','difference_from_original_converted','difference_exceeds_0_1_mpa',
            'candidate_in_original_527','exploratory_label_subset','prior_reliability_status','density_assumption_status','strength_basis_status',
            'not_in_exploratory_subset_reasons','prior_pending_reasons','training_authorized','trusted_qualification_granted']
    for name,rows in [('全部原始记录及暂定标签',data),('探索标签子集_非可信集',eligible),('527候选暂未纳入',[r for r in candidates_out if not r['exploratory_label_subset']])]:
        s=workbook.create_sheet(name);s.append(fields)
        for item in rows:
            flat={**item['original_values'],**item}
            vals=[]
            for f in fields:
                v=flat.get(f)
                if isinstance(v,(list,dict)):v=json.dumps(v,ensure_ascii=False)
                vals.append(v)
            s.append(vals)
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
    s=workbook.create_sheet('用途说明');s.append(['事项','说明'])
    s.append(['用途','探索性暂定强度标签；不是已核实训练集，也不包含已完成预处理的特征矩阵。'])
    s.append(['规则','仅使用各条已有0.95、0.8、0.76；不猜测缺失系数、不替换其他系数。'])
    s.append(['研究范围','子集仍限原527候选与Excel Underwater，ConvertFlag=0且试件字段非占位；不恢复Wang等候选外记录。'])
    s.append(['未知','零值、单位、密度假设、原文成型范围等仍须分开记录；没有自动授予可信资格。'])
    workbook.save(output/'探索标签_非可信训练集.xlsx');workbook.close()
    raw.close();annotated.close();screening.close()
    return summary
