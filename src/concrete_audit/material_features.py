"""独立的材料语义候选层：按证据绑定单位/基准，不改变可信状态或训练输入。"""
from collections import Counter
from datetime import datetime, timezone
import csv
import json
import math
from pathlib import Path
import re
import subprocess
import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

from .field_review import digest, validate_evidence

# 固定字典阻止目标、强度派生信息和论文身份被当成特征。
FEATURES = {
    'material_system': ('category', '材料体系'),
    'antiwashout_product': ('category', '产品/聚合物名称，不表示产品等效'),
    'fiber_type': ('category', '本轮已定位的纤维种类，不推断其他纤维不存在'),
    'polymer_pct_cement': ('%', 'cement_mass'),
    'awa_pct_binder': ('%', 'binder_mass'),
    'awa_mass_kg_m3': ('kg/m3', 'mixture_volume'),
    'steel_mass_ratio_cement': ('kg/kg', 'cement_mass'),
    'pva_vol_pct': ('%', 'mixture_volume'),
    'pva_mass_kg_m3': ('kg/m3', 'mixture_volume'),
    'pva_density_kg_m3': ('kg/m3', 'fiber_volume'),
    'awa_pct_cement': ('%', 'cement_mass'),
    'awa_volume_L_per_100kg_binder': ('L/100kg', 'binder_mass'),
    'pp_mass_kg_m3': ('kg/m3', 'mixture_volume'),
    'sp_volume_L_m3': ('L/m3', 'mixture_volume'),
    'retarder_mL_per_100kg_binder': ('mL/100kg', 'binder_mass'),
    'copolymer_volume_L_m3': ('L/m3', 'mixture_volume'),
    'viscosity_modifier_pct_cement': ('%', 'cement_mass'),
    'sp_pct_cement': ('%', 'cement_mass'),
    'sp_reported_pct_cement': ('%', 'cement_percent_definition_unknown'),
    'awa_reported_pct_cement': ('%', 'cement_percent_definition_unknown'),
    'retarder_pct_cement': ('%', 'cement_mass'),
    'sp_volume_L_per_100kg_cement': ('L/100kg', 'cement_mass'),
    'upa_reported_percent_water_medium': ('%', 'water_medium_percent_definition_unknown'),
    'water_environment': ('category', '外部浇筑水环境'),
    'water_motion_condition': ('category', '外部水流条件'),
    'water_velocity_m_s': ('m/s', 'external_water_velocity'),
    'exposure_protocol': ('category', '暴露/装模程序，入模资格另查'),
}
RAW_FIELDS = ('AWA', 'AWA_Dos', 'AWA_Dos_Base', 'AWA_Type', 'AWA_Form',
              'StF', 'SyF', 'StFiberType', 'SyFiberType', 'SyFiberSG', 'SP', 'Retarder', 'Emulsion',
              'UPA', 'UPA_Type', 'UPA_Conc', 'WaterEnvType', 'WaterVelocity', 'CastMethod')


def append_excel_row(sheet, values):
    """仅Excel展示层去XML禁用字符；JSON原文不变，字符串按文本而非公式保存。"""
    sheet.append([ILLEGAL_CHARACTERS_RE.sub('', x) if isinstance(x, str) else x for x in values])
    for cell in sheet[sheet.max_row]:
        if isinstance(cell.value, str):
            cell.data_type = 's'


def match_source_recipe(specimen_id, recipes):
    """只匹配已从来源表编录的完整配方ID；禁止近似匹配或按论文整批填值。"""
    matches = [r for r in recipes if specimen_id == r['specimen_id']]
    if len(matches) != 1:
        raise ValueError('来源表配方不存在或不唯一')
    return matches[0]


def pva_mass(vol_pct, density_kg_m3):
    """体积百分比÷100×纤维密度；不使用混凝土密度或假设产品密度。"""
    if any(isinstance(x, bool) or not isinstance(x, (int, float)) or not math.isfinite(x)
           for x in (vol_pct, density_kg_m3)):
        raise ValueError('PVA计算需要有限数值')
    if not 0 <= vol_pct <= 100 or density_kg_m3 <= 0:
        raise ValueError('体积百分比或纤维密度无效')
    return vol_pct / 100 * density_kg_m3


def validate_annotation(annotation, pages, pdf_sha):
    """检查固定字段字典、计算链及可定位证据；文本匹配不替代人工审批。"""
    name = annotation['feature']
    if name not in FEATURES:
        raise ValueError('非材料白名单字段')
    if (annotation['unit'], annotation['basis']) != FEATURES[name]:
        raise ValueError('单位或掺量基准错配')
    validate_evidence(annotation, pages, pdf_sha)
    value = annotation['value']
    if annotation['unit'] == 'category':
        if not isinstance(value, str) or not value.strip():
            raise ValueError('分类值为空')
    else:
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value < 0:
            raise ValueError('材料值无效')
        if value == 0:
            zero_ref = annotation.get('explicit_absence_quote', '')
            reported_zero = annotation.get('reported_zero_line', '')
            explicit = zero_ref and any(zero_ref in ref['quote'] for ref in annotation['evidence'])
            # 数值0必须保留完整源表行和列语义，不能只搜索一个字符0。
            dose_tokens = re.findall(r'(?<!\S)\d+(?:\.\d+)?(?!\S)', reported_zero)
            numeric_zero = (name == 'sp_volume_L_m3'
                            and annotation.get('zero_semantics') == 'explicit_reported_dose'
                            and reported_zero.strip()
                            and annotation.get('source_column') == 'High-range water reducing agent L/m3'
                            and any(reported_zero in ref['quote'] for ref in annotation['evidence'])
                            and dose_tokens and float(dose_tokens[0]) == 0)
            upa_zero = (name == 'upa_reported_percent_water_medium'
                        and annotation.get('zero_semantics') == 'explicit_reported_dose'
                        and annotation.get('source_column') == 'UPA (%)'
                        and any(reported_zero in ref['quote'] for ref in annotation['evidence'])
                        and len(dose_tokens) > 1 and float(dose_tokens[1]) == 0)
            if not explicit and not numeric_zero and not upa_zero:
                raise ValueError('零值必须绑定原文明确未掺证据')
    if annotation.get('method') == 'pva_volume_density':
        if name != 'pva_mass_kg_m3':
            raise ValueError('PVA计算不能用于其他字段')
        p = annotation['parameters']
        if not math.isclose(value, pva_mass(p['vol_pct'], p['density_kg_m3']), abs_tol=1e-10):
            raise ValueError('PVA计算链不一致')


def make_views(labels, config, split_rows, pages):
    """仅凭明确记录映射生成候选值；其余None，不按论文名或零值扩散。"""
    by_row = {r['source_row']: r for r in labels}
    if len(by_row) != len(labels) or len({r['sample_id'] for r in labels}) != len(labels):
        raise ValueError('来源行或样本ID重复')
    split = {r['sample_id']: r['split'] for r in split_rows}
    if len(split) != len(split_rows) or set(split) - {r['sample_id'] for r in labels}:
        raise ValueError('划分存在重复或未知样本')
    annotated = {}
    for item in config['annotations']:
        row = item['source_row']
        if row not in by_row or by_row[row]['original_values']['Specimen ID'] != item['specimen_id']:
            raise ValueError('证据与配方身份不匹配')
        if by_row[row]['original_values']['PDF filename'] != item['registry_pdf_filename']:
            raise ValueError('证据与源论文身份不匹配')
        if item.get('source_recipe'):
            recipe = match_source_recipe(item['specimen_id'], config.get('source_recipes', []))
            if recipe['recipe_key'] != item['source_recipe'] or recipe['paper'] != item['paper']:
                raise ValueError('来源配方表绑定不一致')
            if item['feature'] == 'sp_volume_L_m3' and recipe.get('sp_L_m3') != item['value']:
                raise ValueError('来源配方SP剂量不一致')
        key = (row, item['feature'])
        if key in annotated:
            raise ValueError('重复材料字段声明')
        validate_annotation(item, pages[item['paper']], config['papers'][item['paper']]['sha256'])
        annotated[key] = item
    rows = []
    for src in labels:
        row = src['source_row']
        values = {f: annotated[(row, f)]['value'] if (row, f) in annotated else None for f in FEATURES}
        rows.append({
            'sample_id': src['sample_id'], 'source_row': row,
            'split_for_coverage_only': split.get(src['sample_id'], 'core_unknown_not_in_split'),
            'candidate_values': values,
            'unknown_reasons': {f: '本轮未编录适用于该配方和基准的证据；不表示论文未报告' for f, v in values.items() if v is None},
            'raw_recorded_values': {f: src['original_values'].get(f) for f in RAW_FIELDS},
            'prior_reliability_status': src['prior_reliability_status'],
            'source_pdf_for_audit_only': src['original_values']['PDF filename'],
            'human_confirmation': '待核查', 'admitted_to_model': False,
        })
    coverage = {}
    # 测试部分只统计非目标字段覆盖；不访问目标值、预测或评分。
    for scope in ('all_458', 'train', 'test', 'core_unknown_not_in_split'):
        selected = rows if scope == 'all_458' else [r for r in rows if r['split_for_coverage_only'] == scope]
        coverage[scope] = {'denominator': len(selected), 'features': {f: {
            'evidence_candidate_count': sum(r['candidate_values'][f] is not None for r in selected),
            'unknown_count': sum(r['candidate_values'][f] is None for r in selected),
            'human_approved_count': 0,
        } for f in FEATURES}}
    return rows, coverage


def training_readiness(rows, config):
    """只统计训练部分的来源覆盖和缺少论文清单，不接触目标或评分。"""
    training = [r for r in rows if r['split_for_coverage_only'] == 'train']
    available = {a['registry_pdf_filename'] for a in config['annotations']}
    # 仅已在训练集内出现且未有本轮定位编录的来源；缺文件与缺字段分别表述。
    available |= set(config.get('other_available_registry_filenames', []))
    counts = Counter(r['source_pdf_for_audit_only'] for r in training)
    missing = [{'registry_pdf_filename': name, 'training_records': n,
                'reason': '未能与当前已上传原文建立文件对应，不等于证明论文不存在'}
               for name, n in sorted(counts.items(), key=lambda pair: (-pair[1], pair[0])) if name not in available]
    fields = {}
    for feature in FEATURES:
        known = [r for r in training if r['candidate_values'][feature] is not None]
        fields[feature] = {'candidate_records': len(known), 'training_denominator': len(training),
                           'candidate_sources': len({r['source_pdf_for_audit_only'] for r in known}),
                           'complete_on_training': bool(training) and len(known) == len(training)}
    return {'training_records': len(training), 'training_sources': len(counts), 'fields': fields,
            'unmatched_source_pdfs': missing, 'selection_uses_test_or_target': False,
            'automatic_feature_admission': False,
            'decision': ('存在训练全覆盖的候选字段，仍需检查证据与来源泛化再决定入模。'
                         if any(f['complete_on_training'] for f in fields.values()) else
                         '新增材料字段尚无训练全覆盖列；继续补证据，不以未知值补0直接重训。')}


def run(labels_path, config_path, split_path, output):
    """重读PDF校验引文；哈希保护输入，只写全新输出目录，不调用模型。"""
    labels_path, config_path, split_path, output = map(Path, (labels_path, config_path, split_path, output))
    config = json.loads(config_path.read_text())
    protected = [labels_path, config_path, split_path, *map(Path, config['protected_inputs']),
                 *[Path(p['path']) for p in config['papers'].values()]]
    before = {str(p): digest(p) for p in protected}
    pages = {}
    for key, paper in config['papers'].items():
        if digest(paper['path']) != paper['sha256']:
            raise ValueError('PDF版本改变')
        result = subprocess.run(['pdftotext', '-layout', paper['path'], '-'], check=True, capture_output=True)
        pages[key] = result.stdout.decode('utf-8').split('\f')
    labels = json.loads(labels_path.read_text())
    rows, coverage = make_views(labels, config, json.loads(split_path.read_text())['rows'], pages)
    readiness = training_readiness(rows, config)
    if any(digest(p) != h for p, h in before.items()):
        raise RuntimeError('输入内容发生改变')
    evidence_sources = Counter(a['paper'] for a in config['annotations'])
    summary = {'records': len(rows), 'records_with_evidence_candidates': sum(any(v is not None for v in r['candidate_values'].values()) for r in rows),
               'field_annotations': len(config['annotations']), 'new_trusted_samples': 0,
               'annotation_counts_by_paper': dict(evidence_sources),
               'models_trained': False, 'test_performance_evaluated': False, 'inputs_unchanged': True}
    output.mkdir(parents=True, exist_ok=False)
    for name, data in [('candidate_features.json', rows), ('coverage.json', coverage), ('summary.json', summary),
                       ('training_readiness.json', readiness),
                       ('evidence_snapshot.json', config), ('feature_dictionary.json', FEATURES),
                       ('manifest.json', {'inputs_sha256': before, 'module_sha256': digest(__file__),
                                          'created_at_utc': datetime.now(timezone.utc).isoformat()})]:
        (output / name).write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False) + '\n')
    with (output / '材料候选字段.csv').open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=['source_row', 'split', *FEATURES, 'human_confirmation'])
        writer.writeheader()
        for r in rows:
            writer.writerow({'source_row': r['source_row'], 'split': r['split_for_coverage_only'],
                             **r['candidate_values'], 'human_confirmation': r['human_confirmation']})
    # 面向人工审查的独立工作簿：候选值、原文证据、单位字典和训练来源缺口。
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = '候选字段'
    append_excel_row(sheet, ['source_row', 'split', *FEATURES, 'human_confirmation'])
    for r in rows:
        append_excel_row(sheet, [r['source_row'], r['split_for_coverage_only'], *r['candidate_values'].values(), '待核查'])
    evidence_sheet = book.create_sheet('字段证据')
    append_excel_row(evidence_sheet, ['原始行', '配方ID', '候选字段', '候选值', '单位', '基准', '论文', 'PDF页码', '定位', '原文', '说明', '人工状态'])
    for a in config['annotations']:
        append_excel_row(evidence_sheet, [a['source_row'], a['specimen_id'], a['feature'], a['value'], a['unit'], a['basis'], a['paper'],
                               ','.join(str(e['page']) for e in a['evidence']),
                               '\n'.join(e['location'] for e in a['evidence']),
                               '\n\n'.join(e['quote'] for e in a['evidence']), a.get('note', ''), '待核查'])
    dictionary = book.create_sheet('字段与规则')
    append_excel_row(dictionary, ['字段', '单位', '基准/含义'])
    for f, (unit, basis) in FEATURES.items():
        append_excel_row(dictionary, [f, unit, basis])
    append_excel_row(dictionary, ['空白', '未知', '不是0；正值或计算一致也不代表整条已核实可用'])
    append_excel_row(dictionary, ['摘录展示', 'XML控制字符仅在Excel去除', '完整原文保留在evidence_snapshot.json；页码和引文校验仍用原始文本'])
    gap_sheet = book.create_sheet('训练来源缺口')
    append_excel_row(gap_sheet, ['原库论文文件名', '训练记录数', '状态'])
    for item in readiness['unmatched_source_pdfs']:
        append_excel_row(gap_sheet, [item['registry_pdf_filename'], item['training_records'], item['reason']])
    for sheet in book:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
    book.save(output / '材料字段与证据.xlsx')
    lines = ['# 材料语义与证据覆盖报告', '',
             f'范围：现有{len(rows)}条暂定标签探索记录；字段声明涉及{len(evidence_sources)}篇原文，不宣称已穷尽其他原文。',
             '所有候选值均待人工确认，不提升整条可信状态、不加入模型。空值为未知，CSV空白不是0。', '',
             f'|候选字段|单位/基准|有证据候选/{len(rows)}|训练部分候选/{coverage["train"]["denominator"]}|', '|---|---|---:|---:|']
    for f, (unit, basis) in FEATURES.items():
        a = coverage['all_458']['features'][f]['evidence_candidate_count']
        b = coverage['train']['features'][f]['evidence_candidate_count']
        lines.append(f'|{f}|{unit} / {basis}|{a}|{b}|')
    lines += ['', '零值只按明确未掺证据或完整来源表中的明确剂量0逐配方保存；其他未编录记录保留未知。',
              'Zhang绝对kg/m³基准仍未知，按水泥比例保存，不借此生成绝对用量。',
              'Al-Selmi PVA体积百分比与质量分开：1.5/100×1280=19.2、1.75/100×1280=22.4 kg/m³。',
              'PVA密度1.28 g/cm³来自p5 Table3，换算1280 kg/m³；不是假定混凝土密度。',
              'Al-Selmi胶材表注1502与分项1202+301=1503的矛盾保留；不由AWA百分比重算或覆盖表4质量。',
              '不同产品和基准不合并。材料体系标签可能代理来源，未来需按来源分组验证，不能用论文ID代替材料特征。',
              '本轮测试部分只统计输入覆盖，未读取目标进行选择或重新评价。',
              '下一步先扩展既有论文中的同类字段证据，并确定共同可用特征，再另建训练配置；当前覆盖不足以直接加入全部模型。']
    lines += ['', *config.get('report_notes', [])]
    lines += ['', '## 仅训练部分的下一步判断', '',
              f'训练部分{readiness["training_records"]}条，涉及{readiness["training_sources"]}个论文来源；未匹配已上传全文的来源为{len(readiness["unmatched_source_pdfs"])}个。',
              readiness['decision'],
              '缺口清单按训练记录数排列，仅用于补证据优先级，不代表按误差筛样本或修改训练/测试划分。']
    (output / 'report.md').write_text('\n'.join(lines) + '\n')
    return summary
