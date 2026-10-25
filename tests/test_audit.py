"""验证危险边界：算术通过不授信、哈希绑定、缺证据、只读与版本冲突。"""
import tempfile
import unittest
from pathlib import Path
import openpyxl
from concrete_audit.audit import run, sha256, comment_rows
from concrete_audit.reliability import conversion_check, apply_review, ASPECTS, number


class ReliabilityTests(unittest.TestCase):
    def setUp(self):
        self.row = {'CS_28d':20,'Conversion Factor':0.95,'Converted CS_28d':19,
                    'ConvertFlag':0,'CS_Spec':'Cy100'}
        self.record = {'record_id':'x','source_sha256':'hash','casting_env':'Underwater',
                       'status':'待核查','conversion':conversion_check(self.row)}
        self.decision = {'record_id':'x','source_sha256':'hash','reviewer':'核查人',
                         'reviewed_at':'2026-10-21','reason':'逐项核对',
                         'evidence':[{'sha256':'paper','location':'p3 Table 2'}],
                         'status':'已核实可用','checks':{k:'已核实' for k in ASPECTS},
                         'density_assumption':'已核实未使用假设密度反算','unresolved_reasons':[]}

    def test_arithmetic_does_not_verify(self):
        self.assertTrue(self.record['conversion']['passed'])
        self.assertEqual(self.record['status'],'待核查')

    def test_invalid_numbers(self):
        for x in [True,'nan','inf','-','/', '=A1*B1',None]:
            self.assertIsNone(number(x))
        self.assertFalse(conversion_check({**self.row,'Conversion Factor':0})['passed'])
        self.assertFalse(conversion_check({**self.row,'Converted CS_28d':20})['passed'])

    def test_review_gate(self):
        r = apply_review(self.record,self.decision,{'paper'})
        self.assertEqual(r['status'],'已核实可用')
        self.assertEqual(self.record['status'],'待核查')
        for change in [{'source_sha256':'wrong'},{'evidence':[]},{'checks':{}},
                       {'density_assumption':'未知'},{'unresolved_reasons':['待查']},
                       {'reviewer':''},{'record_id':'wrong'}]:
            with self.subTest(change=change), self.assertRaises(ValueError):
                apply_review(self.record,{**self.decision,**change},{'paper'})
        with self.assertRaises(ValueError):
            apply_review(self.record,self.decision,set())

    def test_explicit_range(self):
        self.assertEqual(comment_rows('DF7:DF10；不能把0当真零',7),{7,8,9,10})
        self.assertEqual(comment_rows('某些行已核对',7),{7})


class IntegrationTests(unittest.TestCase):
    def test_readonly_unknown_and_conflicts(self):
        with tempfile.TemporaryDirectory() as temp:
            p=Path(temp); raw=p/'raw.xlsx'; screening=p/'screen.xlsx'
            w=openpyxl.Workbook(); s=w.active; s.title='Data'
            fields=['PDF filename','Research Group','Specimen ID','CastingEnv','CS_28d',
                    'Conversion Factor','Converted CS_28d','ConvertFlag','CS_Spec','Density']
            for r,env in [(7,'Underwater'),(8,'Air')]:
                for i,v in enumerate(['paper.pdf','Group','s'+str(r),env,20,.95,19,0,'Cy100',0],1):
                    s.cell(r,i,v)
            s['E7'].comment=openpyxl.comments.Comment('E7:E8 已核对数值','reviewer')
            d=w.create_sheet('Parameter Explanation'); d.append(['No','Class','Sub','Symbol','Unit','X','Definition'])
            for i,name in enumerate(fields):d.append([i,'','',name,'-',None,'说明'])
            w.save(raw)
            w=openpyxl.Workbook();w.remove(w.active)
            c=w.create_sheet('初筛候选527');c.append(['原表行号',*fields]);c.append([7,'paper.pdf','Group','s7','Underwater',21,.95,19,0,'Cy100',0])
            q=w.create_sheet('待核查清单');q.append(['原表行号',*fields]);q.append([7,'wrong.pdf','Group','s7','Underwater',20,.95,19,0,'Cy100',0])
            f=w.create_sheet('特征完整性');f.append(['字段','候选总数','数值记录数','数值为0记录数','正值记录数','占位/非数值记录数']);f.append(['CS_28d',1,1,0,1,0])
            for name in ['筛选规则','样本统计']:w.create_sheet(name).append(['说明'])
            w.save(screening)
            before=[sha256(raw),sha256(screening)]
            summary=run(raw,screening,p/'out')
            self.assertEqual(summary['status_counts'],{'待核查':1,'已核实可用':0,'已排除':1})
            self.assertEqual(summary['identity_mismatch_links'],1)
            self.assertGreater(summary['cross_version_field_differences'],0)
            self.assertEqual(before,[sha256(raw),sha256(screening)])
            with self.assertRaises(FileExistsError):run(raw,screening,p/'out')


class AdditionalGateTests(unittest.TestCase):
    def test_air_cannot_be_verified(self):
        t=ReliabilityTests();t.setUp()
        t.record['casting_env']='Air'
        with self.assertRaises(ValueError):apply_review(t.record,t.decision,{'paper'})

    def test_unresolved_density_cannot_be_promoted(self):
        t=ReliabilityTests();t.setUp()
        t.decision['checks']['density_assumption']='未知'
        with self.assertRaises(ValueError):apply_review(t.record,t.decision,{'paper'})

if __name__=='__main__':unittest.main()
