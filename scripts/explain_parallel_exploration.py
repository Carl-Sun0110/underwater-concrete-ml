"""训练折重拟合最优主要候选，解释折外预测；不调用历史测试集。"""
import argparse
from collections import defaultdict
import json
from pathlib import Path

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
import shap
from openpyxl import Workbook

from concrete_audit.baselines import grouped_folds
from concrete_audit.material_features import append_excel_row
from concrete_audit.parallel_exploration import digest, feature_matrix, make_pipeline, write_json


def explain(run_dir, output):
    run_dir, output = Path(run_dir), Path(output)
    output.mkdir(parents=True,exist_ok=False)
    manifest=json.loads((run_dir/'manifest.json').read_text())
    # 模型、数据、分组入口必须与实际实验版本相同，避免解释另一批数据。
    for path, expected in manifest['input_hashes_after'].items():
        if digest(path)!=expected:raise ValueError(f'输入版本变化：{path}')
    config=manifest['config']
    results=json.loads((run_dir/'results.json').read_text())
    best=next(r for r in results if r['experiment']==manifest['selected_primary_experiment'])
    source={r['sample_id']:r for r in json.loads(Path(config['labels']).read_text())}
    snapshot=json.loads((run_dir/'quality_snapshot.json').read_text())
    dev_meta=[r for r in snapshot if r['split']=='development' and r['grade']!='D']
    rows=[source[r['sample_id']] for r in dev_meta]
    groups=np.array([r['group'] for r in dev_meta]);mask=np.array([r['grade'] in 'AB' for r in dev_meta])
    y=np.array([r['Provisional_Converted_CS_28d'] for r in rows])
    fields=config['feature_sets'][best['feature_set']];X=feature_matrix(rows,fields)
    folds=grouped_folds(X,y,groups,config['folds'])
    explanations=[];max_residual=0.
    for k,(train,val) in enumerate(folds):
        train,val=train[mask[train]],val[mask[val]]
        if not len(val):continue
        pipe=make_pipeline(best['model'],best['imputation'],config);pipe.fit(X[train],y[train])
        transformed=pipe[:-1].transform(X[val])
        names=pipe['impute'].get_feature_names_out(fields).tolist()
        explainer=shap.TreeExplainer(pipe['model'])
        sv=explainer.shap_values(transformed)
        base=float(np.asarray(explainer.expected_value).reshape(-1)[0])
        pred=pipe.predict(X[val])
        residual=float(np.max(np.abs(base+np.sum(sv,axis=1)-pred)))
        max_residual=max(max_residual,residual)
        if residual>1e-3:raise ValueError('SHAP加和未复现模型预测')
        for j,i in enumerate(val):
            explanations.append({'sample_id':rows[i]['sample_id'],'fold':k,'base_value':base,
                                 'prediction':float(pred[j]),'shap_MPa':dict(zip(names,sv[j].astype(float).tolist()))})
    names=sorted({name for r in explanations for name in r['shap_MPa']})
    mean_abs={name:float(np.mean([abs(r['shap_MPa'].get(name,0.)) for r in explanations])) for name in names}
    write_json(output/'oof_shap.json',{'experiment':best['experiment'],'shap_version':shap.__version__,
        'max_additivity_residual_MPa':max_residual,'mean_absolute_shap_MPa':mean_abs,'rows':explanations,
        'interpretation':'训练折模型对折外预测的关联解释；非因果；缺失指示器可能反映论文记录习惯。折模型不同，基值不相同。'})
    top=sorted(mean_abs,key=mean_abs.get,reverse=True)[:18][::-1]
    fig,ax=plt.subplots(figsize=(8,6));ax.barh(top,[mean_abs[f] for f in top]);ax.set_xlabel('Mean absolute SHAP (MPa), out-of-fold')
    fig.tight_layout();fig.savefig(output/'shap_importance.png',dpi=160);plt.close(fig)
    pred=[r for r in json.loads((run_dir/'oof_predictions.json').read_text()) if r['experiment']==best['experiment']]
    fig,ax=plt.subplots(figsize=(6,6));ax.scatter([r['observed'] for r in pred],[r['predicted'] for r in pred],alpha=.6,s=18)
    low=min(min(r['observed'],r['predicted']) for r in pred);high=max(max(r['observed'],r['predicted']) for r in pred)
    ax.plot([low,high],[low,high],color='black',linestyle='--');ax.set_xlabel('Observed provisional 28d CS (MPa)');ax.set_ylabel('OOF predicted CS (MPa)')
    ax.set_title(f"Grouped CV: MAE={best['MAE_MPa']:.2f}, R2={best['R2']:.3f}")
    fig.tight_layout();fig.savefig(output/'oof_predictions.png',dpi=160);plt.close(fig)
    by_paper=defaultdict(list)
    for r in pred:by_paper[r['paper']].append(r)
    priorities=[]
    total=sum(r['absolute_error'] for r in pred)
    for paper,rr in by_paper.items():
        priorities.append({'paper':paper,'n':len(rr),'MAE_MPa':float(np.mean([r['absolute_error'] for r in rr])),
                           'bias_pred_minus_true_MPa':float(np.mean([r['predicted']-r['observed'] for r in rr])),
                           'total_absolute_error_share':sum(r['absolute_error'] for r in rr)/total,
                           'source_rows':[r['source_row'] for r in rr]})
    priorities.sort(key=lambda r:r['MAE_MPa'],reverse=True)
    write_json(output/'paper_review_priorities.json',priorities)
    perm=json.loads((run_dir/'validation_permutation_importance.json').read_text())
    perm_summary={f:float(np.average([r['MAE_increase_by_feature'][f] for r in perm],
                                    weights=[r['validation_n'] for r in perm])) for f in fields}
    write_json(output/'importance_summary.json',{'permutation_MAE_increase_MPa':perm_summary,
                                               'SHAP_mean_absolute_MPa':mean_abs})
    wb=Workbook();wb.remove(wb.active)
    tables={'实验指标':(['experiment','n','MAE_MPa','RMSE_MPa','R2','macro_group_MAE_MPa'],results),
            '探索分层':(['sample_id','source_row','paper','group','split','grade','reason','prior_reliability_status'],snapshot),
            '主要模型折外预测':(['source_row','paper','group','grade','fold','observed','predicted','absolute_error'],pred),
            '论文核查优先':(['paper','n','MAE_MPa','bias_pred_minus_true_MPa','total_absolute_error_share','source_rows'],priorities)}
    for title,(keys,data) in tables.items():
        sheet=wb.create_sheet(title);append_excel_row(sheet,keys)
        for r in data:append_excel_row(sheet,[json.dumps(r.get(k),ensure_ascii=False) if isinstance(r.get(k),(list,dict)) else r.get(k) for k in keys])
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    notes=wb.create_sheet('使用说明')
    for line in ['全部结果为探索性，不是最终科研结论。','原始标签按用户许可0.95/0.8/0.76暂定换算，不补造标签。',
                 'AB与ABC实验共享来源分组折；质量快照不改变正式待核查状态。','历史145条留出本轮未用于拟合、插补、选模或评估。',
                 'SP/AWA单位未全部统一，仅作敏感性分析。','SHAP为关联解释；CV选模存在选择偏差。']:
        append_excel_row(notes,[line])
    wb.save(output/'探索实验汇总.xlsx')
    write_json(output/'explanation_manifest.json',{'source_manifest_sha256':digest(run_dir/'manifest.json'),
        'script_sha256':digest(__file__),'shap_version':shap.__version__,'n_oof_explained':len(explanations),
        'max_additivity_residual_MPa':max_residual,'input_hashes_rechecked':True})
    (output/'report.md').write_text('# 折外解释与核查重点\n\n'
        +f'对主要候选{best["experiment"]}的{len(explanations)}条折外预测进行Tree SHAP解释；最大加和残差{max_residual:.8f} MPa。\n\n'
        +'高误差论文（不是删除清单）：\n\n'
        +'\n'.join(f'- {r["paper"]}：n={r["n"]}，MAE={r["MAE_MPa"]:.2f} MPa。' for r in priorities[:5])
        +'\n\n优先核查不同材料体系、配合比绝对用量/相对比例及标签换算适用性；已核实的高强度测量值不得为改善指标而删改。SHAP不能证明特征具有因果作用，也不能证明原始数据真实。\n')
    print(json.dumps({'n_explained':len(explanations),'max_additivity_residual':max_residual,'top_SHAP':sorted(mean_abs.items(),key=lambda x:-x[1])[:6]},ensure_ascii=False))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--run',default='outputs/parallel-exploration-v1')
    parser.add_argument('--output',default='outputs/parallel-explanation-v1');args=parser.parse_args()
    explain(args.run,args.output)
