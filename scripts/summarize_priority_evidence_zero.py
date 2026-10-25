"""输出可人工复核的原表单元格证据包，假设敏感性分数不授予可靠资格。"""
import json
from pathlib import Path
import re
import shutil
import numpy as np
from openpyxl import Workbook, load_workbook
from concrete_audit.baselines import regression_metrics
from concrete_audit.material_features import append_excel_row
from concrete_audit.parallel_exploration import digest, write_json


def summarize():
    root=Path('outputs/priority-evidence-zero-v1');m=json.loads((root/'manifest.json').read_text());config=m['config']
    for p,expected in m['input_hashes_after'].items():assert digest(p)==expected
    results=json.loads((root/'sensitivity_results.json').read_text());pred=json.loads((root/'training_oof_predictions.json').read_text())
    packets=json.loads((root/'record_review_packets.json').read_text());papers=json.loads((root/'paper_review_summary.json').read_text())
    cohort=json.loads((root/'training_cohort.json').read_text());network=json.loads((root/'network_attempts.json').read_text())
    old_pred=json.loads((Path(config['refinement'])/'training_oof_predictions.json').read_text())
    baseline_match={};paired=[]
    for r in results:
        pp=[p for p in pred if p['experiment']==r['experiment']]
        assert len(pp)==len(set(p['sample_id'] for p in pp))==cohort['n']
        assert {p['sample_id'] for p in pp}==set(cohort['training_ids'])
        assert not {p['sample_id'] for p in pp}&set(cohort['heldout_ids'])
        scores=regression_metrics([p['actual'] for p in pp],[p['predicted'] for p in pp])
        for key in ['MAE_MPa','RMSE_MPa','R2']:assert abs(scores[key]-r[key])<1e-9
        if r['variant']=='baseline_zero_unknown':
            name=f'exploratory_underwater__{r["feature_set"]}__median__{r["model"]}'
            old={p['sample_id']:p for p in old_pred if p['experiment']==name}
            difference=max(abs(p['predicted']-old[p['sample_id']]['predicted']) for p in pp)
            assert difference<1e-8 and all(p['actual']==old[p['sample_id']]['actual'] for p in pp)
            baseline_match[name]=difference
    for fs in config['feature_sets']:
        for model in config['models']:
            a=next(r for r in results if (r['feature_set'],r['model'],r['variant'])==(fs,model,'baseline_zero_unknown'))
            b=next(r for r in results if (r['feature_set'],r['model'],r['variant'])==(fs,model,'hypothesis_cement_zero_recorded'))
            paired.append({'feature_set':fs,'model':model,'baseline_MAE_MPa':a['MAE_MPa'],'hypothesis_MAE_MPa':b['MAE_MPa'],
                'delta_MAE_MPa':b['MAE_MPa']-a['MAE_MPa'],'baseline_R2':a['R2'],'hypothesis_R2':b['R2'],
                'delta_macro_source_MAE_MPa':b['macro_source_MAE_MPa']-a['macro_source_MAE_MPa']})
    write_json(root/'paired_hypothesis_comparison.json',paired)
    tests_log=Path('/tmp/underwater-priority-evidence-tests.log').read_text();tests=re.search(r'Ran (\d+) tests in ([\d.]+)s\s+OK',tests_log)
    if not tests:raise ValueError('测试未全部通过')
    shutil.copyfile('/tmp/underwater-priority-evidence-tests.log',root/'tests.log')
    shutil.copyfile('/tmp/underwater-priority-evidence-zero.log',root/'run.log')
    wb=Workbook();wb.remove(wb.active)
    tables={'文献与缺口':(['paper','title','doi_as_recorded','baseline_records','baseline_mean_model_groupCV_MAE_MPa','source_rows','pending_flags','corresponding_pdf_available','reason'],papers),
        '逐条单元格核查包':(['source_row','paper','recorded_fields','pending_flags','excel_cell_refs','evidence_level','formal_status_changed'],packets),
        '零值假设配对':(['feature_set','model','baseline_MAE_MPa','hypothesis_MAE_MPa','delta_MAE_MPa','baseline_R2','hypothesis_R2','delta_macro_source_MAE_MPa'],paired),
        '全部训练内结果':(['experiment','n','MAE_MPa','RMSE_MPa','R2','macro_source_MAE_MPa'],results),
        '网络获取记录':(['url','status','exception_type','error','interpretation'],network)}
    for title,(keys,data) in tables.items():
        ws=wb.create_sheet(title);append_excel_row(ws,keys)
        for row in data:append_excel_row(ws,[json.dumps(row.get(k),ensure_ascii=False) if isinstance(row.get(k),(list,dict)) else row.get(k) for k in keys])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    notes=wb.create_sheet('核查边界')
    for text in ['没有下载或读到这4篇新PDF，文献题名/DOI/备注均为原始Excel陈述。',
        '本地工作簿支持的是核查线索，不等于原文验证或正式数据升级。',
        '网络代理403不代表出版商拒绝开放，也不能验证DOI是否正确。',
        '假设实验仅改变训练矩阵中4条记录的Cem数值，不改原表或标签，不自动认定真零。',
        '已生成的预测模型及CSV零值规则未被本轮假设替换。',
        '有助于模型分数不能证明原始数据真实；反之，误差大也不能证明标签错误。']:
        append_excel_row(notes,[text])
    book=root/'优先文献核查包与零值假设.xlsx';wb.save(book)
    check=load_workbook(book,read_only=True,data_only=False);assert not any(c.data_type=='f' for s in check for row in s for c in row);check.close()
    verification={'papers':len(papers),'review_packets':len(packets),'new_pdfs_acquired':0,'unit_tests_passed':int(tests[1]),
        'zero_hypothesis_candidate_rows':cohort['candidate_rows'],'sensitivity_experiments':len(results),
        'training_records_per_experiment':cohort['n'],'prior_pipeline_prediction_difference':baseline_match,
        'heldout_prediction_count':0,'labels_changed':False,'raw_inputs_unchanged':True,'formal_status_changes':0,
        'summary_script_sha256':digest(__file__)}
    write_json(root/'verification.json',verification)
    lines=['# 小批文献核查线索与水泥零值假设','',
        '本轮根据Baseline V1既有折外误差优先整理四篇缺少原文的来源；没有再次使用留出成绩选择模型。全部引用是原表元数据，未冒充PDF核实。','',
        f'实际整理{len(papers)}篇、{len(packets)}条基线记录。网络请求返回代理Tunnel 403，未下载到新论文；不能判断出版商是否开放或DOI是否正确。','',
        '已发现的原表线索：',
        '- Zhang 2023：8条水下基线记录中，747/751/755/759行为Cem0、GGBS550、AlkaliAct189。与水泥为550的A系列并列，提示C系列可能为无水泥碱激发体系；不能仅凭这些列或题名确认真实未掺。',
        '- 751行Water64，其余上述C系列记录Water128；应核对原文表和有效水/激发剂溶液定义，不能按邻行相似性直接改成128。',
        '- Wang 2024：790行CastMethod明确含Mortar、预置骨料骨架灌浆，Paper List备注也称砂浆强度；需确认标签代表砂浆还是整体混凝土。尚不据此删样本或排除论文。',
        '- Horszczaruk 2017：Paper List备注称CastingEnv=Underwater为推定。当前5条记录仍是扩展探索记录，不视为直接水下成型已核实。',
        '- Horszczaruk 2023：12条高影响记录需核对水下修补、压力养护与成型方式及原始强度表；不能由论文题名概括整组条件。','',
        '## 单一假设的训练内对照','',
        '复用Baseline V1.1扩展训练320条、23来源组和同一五折，锁定147条论文留出。仅对指定4条记录测试Cem0照录；其他未知零值规则不动。无参数搜索，无标签改动，不部署假设模型。',
        '|特征|模型|基线MAE MPa|假设MAE MPa|差值|假设R²|','|---|---|---:|---:|---:|---:|']
    for p in paired:lines.append(f'|{p["feature_set"]}|{p["model"]}|{p["baseline_MAE_MPa"]:.3f}|{p["hypothesis_MAE_MPa"]:.3f}|{p["delta_MAE_MPa"]:.3f}|{p["hypothesis_R2"]:.3f}|')
    lines+=['','两特征XGBoost MAE9.118→9.018 MPa，变化较小。分数不能证明真零；方案受既往误差启发且在同一训练CV评估，存在适应性偏差，非最终泛化结果。',
        f'{tests[1]}项测试通过，六组旧流水线逐条预测精确复现，原始输入/标签/正式可靠状态和已发布预测模型均未变。','',
        '## 最优先需要的材料','',
        'Zhongzhe Zhang等，2023，Effect of antiwashout admixture on compressive strength of alkali-activated underwater non-dispersible concrete，Magazine of Concrete Research。原表DOI：10.1680/jmacr.22.00239（本轮无法联网核验）。',
        '优先核查配合比表A/C系列、水泥和矿渣列、激发剂溶液组成与水量定义、原始28天强度表及150mm试件。取得这一篇原文即可先解决当前4条水泥零值的关键含义，无须等待全部论文。',
        '其余三篇题名、DOI照录、需要核查的记录和单元格位置见Excel与paper_review_summary.json。','',
        '新增文件：configs/priority_evidence_zero_v1.json、src/concrete_audit/evidence_batch.py、scripts/run_priority_evidence_zero.py、scripts/summarize_priority_evidence_zero.py、tests/test_evidence_batch.py；README增加本轮说明。',
        '函数说明：flags只标注原表待核查线索；request_metadata真实记录请求或失败；run绑定原表单元格并执行单一训练内假设；summarize复算指标并检查旧流水线、留出隔离及原始哈希。']
    (root/'report.md').write_text('\n'.join(lines)+'\n')
    print(json.dumps({k:v for k,v in verification.items() if k!='prior_pipeline_prediction_difference'},ensure_ascii=False))


if __name__=='__main__':summarize()
