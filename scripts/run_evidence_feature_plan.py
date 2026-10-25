"""实际生成原文细骨料比、折内预处理诊断和扩展输入方案；不训练模型。"""
from pathlib import Path
import argparse
import json
import subprocess
import sys
import numpy as np
import openpyxl
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import StandardScaler

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
from concrete_audit.field_review import digest
from concrete_audit.evidence_features import (build_feature_records, EvidenceFeatureTransformer,
                                             fold_support, RATIO_FEATURE)
from concrete_audit.material_features import append_excel_row


def run(output):
    output = Path(output)
    config_path = ROOT / 'configs/fine_aggregate_ratio_evidence.json'
    labels_path = ROOT / 'outputs/provisional-labels-v1/exploratory_labels.json'
    scope = ROOT / 'outputs/casting-scope-gate-v1'
    split_path = ROOT / 'outputs/exploratory-baselines-v2/split_manifest.json'
    material_path = ROOT / 'outputs/material-feature-three-papers-v1/candidate_features.json'
    config = json.loads(config_path.read_text())
    views = json.loads((scope / 'scope_views.json').read_text())['strict_evidence_scope']
    expected = views['train_sample_ids']
    test_ids = set(views['protected_original_test_ids'])
    folds = json.loads((scope / 'training_plan.json').read_text())['plans']['strict_evidence_scope']['cv']
    group_by_id = {r['sample_id']: r['leakage_group'] for r in json.loads(split_path.read_text())['rows']}
    protected = [config_path, labels_path, split_path, material_path, scope / 'scope_registry.json',
                 scope / 'scope_views.json', scope / 'training_plan.json',
                 *map(Path, config['protected_inputs']), *[Path(p['path']) for p in config['papers'].values()]]
    # 沿已有范围档案继续保护原始数据和可靠状态，而非只保护派生表。
    upstream = json.loads((scope / 'scope_registry.json').read_text())['protected_inputs_sha256']
    if any(digest(path) != h for path, h in upstream.items()):
        raise ValueError('范围审查上游输入已变化')
    before = {**upstream, **{str(p.resolve()): digest(p) for p in protected}}
    pages = {}
    for key, paper in config['papers'].items():
        if digest(paper['path']) != paper['sha256']:
            raise ValueError('论文版本不符')
        result = subprocess.run(['pdftotext', '-layout', paper['path'], '-'], check=True, capture_output=True)
        pages[key] = result.stdout.decode().split('\f')
    records = build_feature_records(json.loads(labels_path.read_text()), config['annotations'], expected, test_ids, pages, config['papers'])
    by_id = {r['sample_id']: r for r in records}
    X = np.array([[by_id[sid]['candidate_features'][f] for f in ['Water', 'Cem', RATIO_FEATURE]] for sid in expected])
    audits = []
    for f in folds:
        valid_ids = set(f['validation_sample_ids'])
        tr = [i for i, sid in enumerate(expected) if sid not in valid_ids]
        va = [i for i, sid in enumerate(expected) if sid in valid_ids]
        if {group_by_id[expected[i]] for i in tr} & {group_by_id[expected[i]] for i in va}:
            raise ValueError('预处理折存在来源重叠')
        # 真正执行每个外折的fit/transform；不接收y、不创建回归器。
        pipe = Pipeline([('physical_inputs', EvidenceFeatureTransformer(add_water_cement_ratio=True)), ('scale', StandardScaler())])
        pipe.fit(X[tr])
        before_mean = pipe['scale'].mean_.copy()
        transformed = pipe.transform(X[va])
        assert np.isfinite(transformed).all() and np.array_equal(before_mean, pipe['scale'].mean_)
        val_ratio = X[va, 2]
        train_ratio = X[tr, 2]
        audits.append({'fold': f['fold'], 'train_records': len(tr), 'validation_records': len(va),
                       'feature_order': ['Water', 'Cem', RATIO_FEATURE, 'computed_Water_over_Cem'],
                       'scaler_training_mean': pipe['scale'].mean_.tolist(), 'scaler_training_scale': pipe['scale'].scale_.tolist(),
                       'ratio_validation_outside_train_range': int(((val_ratio < train_ratio.min()) | (val_ratio > train_ratio.max())).sum()),
                       'validation_changed_training_statistics': False, 'test_rows_used': 0,
                       'validation_sample_ids': [expected[i] for i in va]})
    material = json.loads(material_path.read_text())
    fields = ['water_environment', 'water_motion_condition', 'water_velocity_m_s', 'material_system',
              'upa_reported_percent_water_medium', 'sp_volume_L_per_100kg_cement', 'awa_mass_kg_m3', 'polymer_pct_cement', 'pva_vol_pct']
    coverage = fold_support(material, expected, group_by_id, folds, fields)
    summary = {'strict_training_records': len(records), 'training_groups': len({group_by_id[sid] for sid in expected}),
               'new_ratio_evidence_coverage': len(records), 'outer_fold_preprocessing_runs': len(audits),
               'new_feature': RATIO_FEATURE, 'imputation_fitted': False, 'models_trained': False,
               'test_rows_used_for_preprocessing': 0, 'new_trusted_samples': 0,
               'inputs_unchanged': all(digest(path) == h for path, h in before.items())}
    if not summary['inputs_unchanged']:
        raise RuntimeError('输入改变')
    output.mkdir(parents=True, exist_ok=False)
    policy = {'stage': '证据特征与预处理预案，不是训练配置',
              'base_features': ['Water', 'Cem'], 'candidate_addition': RATIO_FEATURE,
              'candidate_feature_sets': [['Water', 'Cem', RATIO_FEATURE], ['Water', 'Cem', RATIO_FEATURE, 'computed_Water_over_Cem']],
              'statistical_imputation': None, 'validation_protocol': '既定5外折×4内折；后续模型与缩放在各内训练折重新拟合',
              'reference_comparison': '相同128条、相同外折，与已有核心输入结果作预先指定的逐折比较',
              'model_training_authorized_by_this_file': False,
              'implementation_needed': '现有训练入口仅接受两核心输入；新增证据特征还需独立版本的门禁接入，不能直接修改旧配置绕过。'}
    for name, data in [('candidate_training_features.json', records), ('fold_preprocessing_audit.json', audits),
                       ('other_feature_fold_support.json', coverage), ('feature_preparation_policy.json', policy),
                       ('evidence_snapshot.json', config), ('summary.json', summary),
                       ('manifest.json', {'inputs_sha256': before, 'module_sha256': digest(ROOT / 'src/concrete_audit/evidence_features.py'), 'script_sha256': digest(__file__)})]:
        (output / name).write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False) + '\n')
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = '128条候选输入'
    append_excel_row(sheet, ['原始行', 'Water', 'Cem', RATIO_FEATURE, '人工状态', '新增可信资格'])
    for r in records:
        append_excel_row(sheet, [r['source_row'], *r['candidate_features'].values(), '待核查', '否'])
    sheet = book.create_sheet('原文质量比证据')
    append_excel_row(sheet, ['原始行', '配方', '来源', '细骨料量', '水泥量', '原单位', 'PDF页码', '说明', '原文'])
    for a in config['annotations']:
        append_excel_row(sheet, [a['source_row'], a['specimen_id'], a['paper'], a['numerator'], a['denominator'], a['numerator_unit'],
                                 ','.join(str(e['page']) for e in a['evidence']), a['note'], '\n\n'.join(e['quote'] for e in a['evidence'])])
    sheet = book.create_sheet('其他特征覆盖')
    append_excel_row(sheet, ['字段', '训练有值/128', '来源组数/5', '整折无训练证据', '当前处理'])
    for row in coverage:
        append_excel_row(sheet, [row['field'], row['known_training_records'], row['known_groups'],
                                 ','.join(str(f['fold']) for f in row['fold_checks'] if f['held_out_only_evidence']), row['status']])
    for sheet in book:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
    book.save(output / '新增物理特征与预处理核查.xlsx')
    lines = ['# 新增证据特征与预处理检查', '',
             '本轮新增细骨料/水泥质量比候选，原文证据覆盖严格范围128条训练记录、5个来源组。不是仅对原有Water/Cem重新组合，而是引入原文细骨料配比信息。所有候选仍待人工确认，没有改变整条可靠状态。', '',
             'Zhang以归一化Sand/Cement=1.05直接取得；其他来源以同表kg/m³细骨料与水泥相除。分母是水泥，不是胶凝材料。Yousri的F.A.是细骨料，不是FA粉煤灰；Brzozowski没有使用相邻Water/SP体积列。无需假定混凝土密度，也不据此宣称Zhang原库绝对kg/m³已验证。', '',
             '实际执行5次外折预处理：只在外训练折拟合StandardScaler，验证折仅transform，未接收目标y。没有填补未知值，没有回归模型训练，也未读取测试行用于预处理。该缩放审计不等同于未来内折缩放；未来调参必须在每个内训练折重新拟合。', '',
             '|其他候选字段|有值/128|来源组/5|仅验证折有证据的外折|', '|---|---:|---:|---|']
    for row in coverage:
        blocked = ','.join(str(f['fold']) for f in row['fold_checks'] if f['held_out_only_evidence']) or '无'
        lines.append(f'|{row["field"]}|{row["known_training_records"]}|{row["known_groups"]}|{blocked}|')
    lines += ['', '水环境、流速、UPA浓度等看似覆盖较高，实际集中在Nasr同一组；该组被留出时，训练折没有此列原文证据。不能利用被留出的数据决定填补值或编码，也不能把未知填0冒充观测。',
              'UPA和SP仍有原文浓度/单位疑点，本轮未因覆盖率较高直接准许入模。材料体系类别也可能与来源混杂，仍需独立处理。', '',
              '## 下一轮预案', '',
              '先保持相同128条、相同外折与既有参数网格，只比较加入细骨料质量比前后的训练内折外结果。保留不加计算水灰比和加入计算比值两种表示。其他稀疏/争议列继续独立审计，不混入主比较。',
              '现有模型训练入口只接收Water/Cem，本轮输出是特征准备政策而非可执行训练配置；下一阶段需为证据特征建立独立接入版本，禁止直接改旧文件冒充原方案。', '',
              '关键函数：positive_table_columns精确读取所需原表列；mass_ratio检查同单位/同归一化基准；build_feature_records只生成固定训练ID的X；EvidenceFeatureTransformer检查三列正数并可计算Water/Cem；fold_support检查来源与外折证据覆盖。', '',
              '未解决：原核心Water/Cem和暂定强度仍有历史可靠性疑点；新质量比有原文依据但未人工签字，来源仍只有5组。', '',
              '输出：candidate_training_features.json、evidence_snapshot.json、fold_preprocessing_audit.json、other_feature_fold_support.json和独立Excel。']
    (output / 'report.md').write_text('\n'.join(lines) + '\n')
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    run(parser.parse_args().output)
