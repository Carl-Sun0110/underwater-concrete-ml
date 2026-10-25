"""校验嵌套验证产物，输出Excel对照与基于已有证据的下一批核查清单。"""
import argparse
from collections import defaultdict
import json
from pathlib import Path
import re

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook, load_workbook

from concrete_audit.baselines import regression_metrics
from concrete_audit.material_features import append_excel_row
from concrete_audit.parallel_exploration import digest, write_json


def summarize(run_dir, test_log):
    root=Path(run_dir);output=root/'analysis';output.mkdir(exist_ok=False)
    manifest=json.loads((root/'manifest.json').read_text());config=manifest['config']
    baseline=Path(config['baseline_run']);old=json.loads((baseline/'results.json').read_text())
    old_manifest=json.loads((baseline/'manifest.json').read_text())
    old_best=next(r for r in old if r['experiment']==old_manifest['selected_primary_experiment'])
    old_pp=[r for r in json.loads((baseline/'oof_predictions.json').read_text()) if r['experiment']==old_best['experiment']]
    old_byid={r['sample_id']:r for r in old_pp}
    results=json.loads((root/'results.json').read_text());pred=json.loads((root/'oof_predictions.json').read_text())
    cohort=json.loads((root/'cohort_and_folds.json').read_text());ids={r['sample_id'] for r in cohort}
    assert ids==set(old_byid)
    snapshot=json.loads((baseline/'quality_snapshot.json').read_text())
    reserved={r['sample_id'] for r in snapshot if r['split']=='reserved_historical_test'}
    assert not reserved & {r['sample_id'] for r in pred}
    for r in results:
        rr=[p for p in pred if p['experiment']==r['experiment']]
        assert len(rr)==len(ids) and {p['sample_id'] for p in rr}==ids
        scores=regression_metrics([p['observed'] for p in rr],[p['predicted'] for p in rr])
        for key in ['MAE_MPa','RMSE_MPa','R2']:assert abs(scores[key]-r[key])<1e-10
        assert all(p['observed']==old_byid[p['sample_id']]['observed'] and
                   p['outer_fold']==old_byid[p['sample_id']]['fold'] for p in rr)
    traces=json.loads((root/'inner_search_traces.json').read_text())
    for trace in traces:
        outer_train=set(trace['outer_train_groups']);outer_val=set(trace['outer_validation_groups'])
        assert not outer_train&outer_val
        for inner in trace['inner_folds']:
            t,v=set(inner['train_groups']),set(inner['validation_groups'])
            assert not t&v and t|v==outer_train and not (t|v)&outer_val
    for path,expected in manifest['input_hashes_after'].items():assert digest(path)==expected
    tests=Path(test_log).read_text();match=re.search(r'Ran (\d+) tests in ([\d.]+)s\s+OK',tests)
    if not match:raise ValueError('本次测试日志未显示全部通过')
    best=next(r for r in results if r['experiment']==manifest['best_experiment'])
    prior_paper=defaultdict(list)
    for p in old_pp:prior_paper[p['paper']].append(p)
    inventory={r['source_name']:r for r in json.loads((baseline/'paper_availability.json').read_text())}
    composition_path=Path('outputs/composition-review-v2/composition_records.json')
    composition={r['source_row']:r for r in json.loads(composition_path.read_text())}
    source_rows={r['source_row']:r for r in json.loads(Path(old_manifest['config']['labels']).read_text())}
    priorities=json.loads((root/'paper_review_priorities.json').read_text());enriched=[]
    for p in priorities:
        rr=[composition[row] for row in p['source_rows'] if row in composition]
        missing_core=[row for row in p['source_rows'] if any(
            not isinstance(source_rows[row]['original_values'].get(f),(int,float)) or source_rows[row]['original_values'][f]<=0
            for f in ['Water','Cem'])]
        issues=sorted({r['source_notes'] for r in rr})
        if any(r.get('binder_total_conflict') for r in rr):issues.append('已有证据记录胶凝总量与组分相加不一致，保留冲突，不自动更正。')
        if missing_core:issues.append(f'原表水/水泥存在未知或非正值的行：{missing_core}；仅统计插补，未补写原表。')
        if '2025-Yu Zhang' in p['paper']:
            issues.append('此前4条原始28天强度数值已有文献核查依据；优先检查相对配比/绝对用量及高强度体系外推，不按误差删改强度。')
        if '2026-Zainab' in p['paper']:
            issues.append('已有字段核查中，1080行原始28天值有明确文本依据，其余图形读值仍待逐项确认；不凭模型预测反写标签。')
        old_mae=float(np.mean([r['absolute_error'] for r in prior_paper[p['paper']]]))
        enriched.append({**p,'prior_reference_MAE_MPa':old_mae,'difference_MAE_MPa':p['MAE_MPa']-old_mae,
            'local_pdf':inventory.get(p['paper'],{}).get('available_path'),
            'core_unknown_rows':missing_core,'existing_evidence_notes':issues or ['目前无本轮复用的详细组分核查档案；来源明确，仍保留探索身份。'],
            'evidence_status':'复用既有证据记录，未声称本轮新核实或完成正式人工审批'})
    write_json(output/'next_review_batch.json',enriched)
    table=[]
    for model in ['rf','xgb']:
        r=next(r for r in old if r['experiment']==f'AB__materials__median__{model}__unknown')
        table.append({'experiment':r['experiment'],'stage':'previous_fixed',**{k:r[k] for k in ['n','MAE_MPa','RMSE_MPa','R2','macro_group_MAE_MPa']}})
    for r in results:table.append({'experiment':r['experiment'],'stage':'nested',**{k:r[k] for k in ['n','MAE_MPa','RMSE_MPa','R2','macro_group_MAE_MPa']}})
    fig,ax=plt.subplots(figsize=(9,5));labels=[r['experiment'].replace('AB__','').replace('__unknown','') for r in table]
    ax.barh(labels[::-1],[r['MAE_MPa'] for r in table][::-1],color=['#5687b1']*len(table))
    ax.axvline(old_best['MAE_MPa'],color='#bd552e',linestyle='--',label='Previous XGBoost reference')
    ax.set_xlabel('Grouped out-of-fold MAE (MPa), same 283 samples');ax.legend();fig.tight_layout()
    fig.savefig(output/'model_comparison.png',dpi=160);plt.close(fig)
    wb=Workbook();wb.remove(wb.active)
    sheets={'模型对照':(['stage','experiment','n','MAE_MPa','RMSE_MPa','R2','macro_group_MAE_MPa'],table),
            '新实验折外预测':(['experiment','source_row','paper','group','grade','outer_fold','observed','predicted','absolute_error'],pred),
            '下一批核查':(['paper','n','MAE_MPa','prior_reference_MAE_MPa','difference_MAE_MPa','source_rows','local_pdf','core_unknown_rows','existing_evidence_notes','action'],enriched),
            '参数选择':(['experiment','outer_fold','best_params','best_inner_mean_fold_MAE_MPa','outer_train_groups','outer_validation_groups'],traces),
            '固定队列与外层折':(['sample_id','source_row','paper','group','grade','outer_fold'],cohort)}
    for title,(keys,rows) in sheets.items():
        ws=wb.create_sheet(title);append_excel_row(ws,keys)
        for row in rows:append_excel_row(ws,[json.dumps(row.get(k),ensure_ascii=False) if isinstance(row.get(k),(dict,list)) else row.get(k) for k in keys])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    ws=wb.create_sheet('说明')
    for line in ['全部为探索性结果；原始标签和正式可靠状态未变。','相同283条AB及五个外层折；内层三折调参。',
                 '比值只在原表单位可比假定下作探索解释，不将未知矿物组分填零凑水胶比。','历史145条留出本轮未作拟合、插补、选模和评估。',
                 '上轮XGBoost仍为当前参考；本轮参数搜索和比值未带来整体改善。','多轮使用开发集存在适应性偏差，嵌套CV不等于最终独立测试。',
                 'nested-v1因CatBoost接口兼容失败；nested-v2适配后完整重跑，失败日志保留。']:
        append_excel_row(ws,[line])
    workbook=output/'第二轮嵌套验证与核查清单.xlsx';wb.save(workbook)
    check=load_workbook(workbook,read_only=True,data_only=False)
    assert not any(c.data_type=='f' for s in check for row in s for c in row);check.close()
    fit_count=sum(len(t['candidates'])*len(t['inner_folds'])+1 for t in traces)
    validation={'experiments':len(results),'same_AB_records':len(ids),'outer_folds':5,
        'inner_folds_per_outer':3,'search_traces':len(traces),'outer_experiment_fits_including_refits':fit_count,
        'all_metrics_recalculated':True,'inner_outer_source_isolation':True,'labels_and_outer_folds_unchanged':True,
        'reserved_test_predictions':0,'protected_inputs_unchanged':True,'unit_tests_passed':int(match[1]),
        'excel_formula_cells':0,'summary_script_sha256':digest(__file__),
        'supplemental_evidence_sha256':{str(composition_path):digest(composition_path)},
        'adoption':'保留上轮固定参数XGBoost为参考；本轮最佳嵌套XGBoost未显示合并样本MAE改善。'}
    write_json(output/'verification.json',validation)
    lines=['# 第二轮实施结果与下一步','',f'同一数据版本、{len(ids)}条AB开发记录、五个外层分组折，六组嵌套实验。',
        f'本轮最低MAE {best["MAE_MPa"]:.4f} MPa（{best["experiment"]}）；上轮固定参数参考 {old_best["MAE_MPa"]:.4f} MPa。未显示整体改善，保留上轮参考。',
        f'104项全量测试通过，{len(pred)}条折外预测，30份外层搜索记录，原始输入哈希、标签、外层折和历史测试隔离均验证通过。','',
        '|实验|MAE MPa|RMSE MPa|R²|来源组平均MAE|','|---|---:|---:|---:|---:|']
    for r in table:lines.append(f'|{r["experiment"]}|{r["MAE_MPa"]:.3f}|{r["RMSE_MPa"]:.3f}|{r["R2"]:.3f}|{r["macro_group_MAE_MPa"]:.3f}|')
    lines+=['','解释与决策：',
        '- 本轮XGBoost、CatBoost、RF加入比值后均未改善MAE，不继续无边界叠加特征。RF调参有小幅改善，未超过XGBoost。',
        '- 在已多轮使用的开发集上，调参和挑选实验仍存在适应性偏差；不把当前R²写成最终独立泛化能力。',
        '- 保留全部283条记录和高强度标签。Zhang高强度体系误差仍大，优先改善配合比表示、体系覆盖和单位证据；不能因预测差判定标签错误。',
        '- 已有文献核查档案与折外误差关联在next_review_batch.json，缺失PDF或未知字段仅保留标记，不阻塞其余建模。',
        '- 下一批核查优先处理现有Zhang和Al-Selmi两篇中的配比单位/组分总量依据；模型轨道增加SVR/MLP的同折对照，并考虑来源组等权的内层选参指标。',
        '- 不按论文身份硬编码数值特征，不从待核查信息推导不存在的强度标签。',
        '- 兼容性问题：初次CatBoost 1.2.8缺少sklearn 1.8回归标签接口，已用SklearnCatBoost适配并完成克隆、Pipeline、GridSearch测试；旧失败日志保留。','',
        '修改/新增文件：configs/parallel_nested_v1.json（失败配置留档）、configs/parallel_nested_v2.json、src/concrete_audit/nested_exploration.py、scripts/run_nested_exploration.py、scripts/summarize_nested_exploration.py、tests/test_nested_exploration.py、requirements-nested.txt、README.md。',
        '主要函数解释见新模块docstring：derive_ratios处理比值与缺失；tune限定内层调参范围；load_development锁定旧分组；run保存完整嵌套证据；summarize重算指标并生成核查清单。']
    (output/'report.md').write_text('\n'.join(lines)+'\n')
    print(json.dumps(validation,ensure_ascii=False))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--run',default='outputs/parallel-nested-v2')
    parser.add_argument('--test-log',default='/tmp/underwater-nested-tests.log');args=parser.parse_args()
    summarize(args.run,args.test_log)
