"""校验训练集内部对照与CSV预测演示，明确这些分数不是新留出测试。"""
import csv
import json
from pathlib import Path
import re
import shutil
import joblib
import numpy as np
from openpyxl import Workbook, load_workbook

from concrete_audit.baselines import regression_metrics
from concrete_audit.material_features import append_excel_row
from concrete_audit.parallel_exploration import digest, write_json


def summarize():
    root=Path('outputs/baseline_v1_1');manifest=json.loads((root/'manifest.json').read_text())
    results=json.loads((root/'results.json').read_text());pred=json.loads((root/'training_oof_predictions.json').read_text())
    cohorts=json.loads((root/'cohort_and_training_folds.json').read_text());selected=json.loads((root/'selected_models.json').read_text())
    for p,expected in manifest['input_hashes_after'].items():assert digest(p)==expected
    heldout_bycohort={c['cohort']:set(c['heldout_ids']) for c in cohorts}
    for r in results:
        pp=[p for p in pred if p['experiment']==r['experiment']]
        assert len(pp)==r['n']==len({p['sample_id'] for p in pp})
        assert not set(p['sample_id'] for p in pp)&heldout_bycohort[r['cohort']]
        scores=regression_metrics([p['actual'] for p in pp],[p['predicted'] for p in pp])
        for field in ['MAE_MPa','RMSE_MPa','R2']:assert abs(scores[field]-r[field])<1e-9
    for c in cohorts:
        ids=set(c['train_ids']);validation=[]
        for fold in c['folds']:
            assert not set(fold['train_groups'])&set(fold['validation_groups'])
            assert not set(fold['train_ids'])&set(fold['validation_ids'])
            assert set(fold['train_ids'])|set(fold['validation_ids'])==ids
            validation += fold['validation_ids']
        assert len(validation)==len(ids) and set(validation)==ids
    comparisons=[]
    index={(r['cohort'],r['feature_set'],r['imputation'],r['model']):r for r in results}
    for cohort in heldout_bycohort:
        for model in manifest['config']['models']:
            for fs in manifest['config']['features']:
                a=index[cohort,fs,'median',model];b=index[cohort,fs,'scaled_knn',model]
                comparisons.append({'cohort':cohort,'feature_set':fs,'model':model,'median_MAE_MPa':a['MAE_MPa'],
                    'KNN_MAE_MPa':b['MAE_MPa'],'delta_MAE_MPa':b['MAE_MPa']-a['MAE_MPa'],
                    'delta_macro_source_MAE_MPa':b['macro_source_MAE_MPa']-a['macro_source_MAE_MPa']})
    write_json(root/'paired_imputation_comparison.json',comparisons)
    demo_count=0;models={}
    for s in selected:
        eligible=[r for r in results if r['cohort']==s['cohort'] and r['model']!='mean']
        assert s['experiment']==min(eligible,key=lambda r:r[manifest['config']['selection_metric']])['experiment']
        model=joblib.load(s['model_file']);c=next(c for c in cohorts if c['cohort']==s['cohort'])
        assert model['training_n']==c['train_n'] and not model['heldout_evaluated_this_round']
        input_file=root/f'{s["cohort"]}__demo_inputs.csv';output_file=root/f'{s["cohort"]}__demo_predictions.csv'
        with input_file.open(encoding='utf-8-sig') as f:demo=list(csv.DictReader(f))
        with output_file.open(encoding='utf-8-sig') as f:output=list(csv.DictReader(f))
        X=np.array([[float(r[field]) if r[field] else np.nan for field in model['features']] for r in demo])
        recomputed=model['pipeline'].predict(X)
        np.testing.assert_allclose(recomputed,[float(r['predicted_Converted_CS_28d_MPa']) for r in output],rtol=0,atol=1e-8)
        demo_count+=len(output);models[s['cohort']]={'model_sha256':digest(s['model_file']),
            'demo_input_sha256':digest(input_file),'demo_output_sha256':digest(output_file),'heldout_evaluated':False}
    log=Path('/tmp/underwater-baseline-v1-1-tests.log').read_text();tests=re.search(r'Ran (\d+) tests in ([\d.]+)s\s+OK',log)
    if not tests:raise ValueError('全量测试未通过')
    shutil.copyfile('/tmp/underwater-baseline-v1-1-tests.log',root/'tests.log')
    shutil.copyfile('/tmp/underwater-baseline-v1-1.log',root/'console.log')
    wb=Workbook();wb.remove(wb.active)
    tables={'训练CV全部方案':(['cohort','feature_set','imputation','model','n','cv_folds','MAE_MPa','RMSE_MPa','R2','macro_source_MAE_MPa'],results),
        '插补配对对照':(['cohort','feature_set','model','median_MAE_MPa','KNN_MAE_MPa','delta_MAE_MPa','delta_macro_source_MAE_MPa'],comparisons),
        '选择的探索模型':(['cohort','experiment','training_n','heldout_n','MAE_MPa','RMSE_MPa','R2','macro_source_MAE_MPa','model_file'],selected),
        '训练折外预测':(['experiment','source_row','paper','group','fold','actual','predicted','absolute_error'],pred)}
    for title,(keys,rows) in tables.items():
        ws=wb.create_sheet(title);append_excel_row(ws,keys)
        for row in rows:append_excel_row(ws,[row.get(k) for k in keys])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    notes=wb.create_sheet('使用边界')
    for line in ['分数为固定留出之后的训练集内部CV，不是Baseline V1相同协议分数。',
        '论文留出95/147条本轮不做新的预测或评估。',
        '按训练CV来源组平均MAE选模型；选模与报告同一CV仍可能乐观。',
        '主集只有4个训练来源组，结果很不稳定；R²为负，不能宣称有效泛化。',
        'CSV示例来自训练记录，仅演示流程；输入范围检查不是可靠性证明，也不提供置信区间。',
        '物理单位与原库名义基准一致，不自动转换水体积或密度；所有预测为探索性质。']:
        append_excel_row(notes,[line])
    book=root/'Baseline_V1_1训练内对照.xlsx';wb.save(book)
    check=load_workbook(book,read_only=True,data_only=False);assert not any(c.data_type=='f' for s in check for row in s for c in row);check.close()
    verification={'experiments_including_mean':len(results),'training_oof_predictions':len(pred),
        'unit_tests_passed':int(tests[1]),'heldout_prediction_count':0,'heldout_used_for_selection':False,
        'source_group_isolation':True,'protected_inputs_unchanged':True,'all_metrics_recomputed':True,
        'demo_rows':demo_count,'demo_is_existing_training_data':True,'models':models,
        'summary_script_sha256':digest(__file__)}
    write_json(root/'verification.json',verification)
    lines=['# Baseline V1.1：训练集内缺失处理对照及CSV预测入口','',
        '本轮锁定Baseline V1的论文留出划分；未用留出目标重新评价、选方案或拟合最终模型。分组身份仅从既有输出提取，目标仅取训练ID对应记录。',
        '已用同一训练CV挑选方案和报告，因此仍有选择偏差。历史留出之前已评价过，本轮没有新独立最终测试结果。','',
        '|队列|训练数|训练来源组|CV折数|锁定留出数|','|---|---:|---:|---:|---:|']
    for c in cohorts:lines.append(f'|{c["cohort"]}|{c["train_n"]}|{c["training_source_groups"]}|{c["cv_folds"]}|{c["heldout_n"]}|')
    lines+=['','固定RF/XGBoost/CatBoost参数不变，比较Water/Cem两特征和完整11配合比特征；中位数或标准化距离KNN插补附缺失标记，全部只在训练折拟合。KNN用训练观察值估计距离尺度，填补后还原单位，全未知列保留数值占位。','',
        '|按来源组平均MAE选出的候选|MAE MPa|RMSE MPa|R²|来源组平均MAE|','|---|---:|---:|---:|---:|']
    for s in selected:lines.append(f'|{s["experiment"]}|{s["MAE_MPa"]:.3f}|{s["RMSE_MPa"]:.3f}|{s["R2"]:.3f}|{s["macro_source_MAE_MPa"]:.3f}|')
    lines+=['','结果解释：',
        '- 主集只剩4个训练来源组。Nasr组104条占117条的绝大部分；留出该组做训练CV时，仅剩其他体系13条，造成严重跨域外推。不能把本轮MAE与前轮212条全体CV直接比较，也不能把变化解释为模型代码退化。',
        '- 扩展训练集按预先指定的来源宏平均误差选择两特征/中位数XGBoost。KNN改善部分全特征组合，但没有形成稳定高精度的跨来源模型。',
        '- 两个模型都保留为探索演示；主集R²为负，不能作为有效泛化结论。扩展模型也只略高于均值参照，不能宣称可以准确设计配合比。',
        f'- {tests[1]}项测试通过，两个模型已保存并各生成3条CSV演示预测。演示来自训练样本，不计为新实验数据或独立性能。',
        '- 优先扩充独立来源与真实单位证据；继续保留原样本和强度标签，不以删难样本或重复看测试集来制造提升。','',
        '## 直接预测CSV','',
        'CSV使用模型白名单中的原始字段名。扩展模型此次只需Water、Cem；名义单位遵循原库，单位疑问仍保留。留空或非正值会标记未知并使用训练时拟合的填补器；多余列（包括强度标签）不作为输入。',
        '```bash',
        'PYTHONPATH=src /workspace/.venv-underwater-ml/bin/python scripts/predict_strength.py --model outputs/baseline_v1_1/exploratory_underwater__selected.joblib --input your_recipes.csv --output your_predictions.csv',
        '```','',
        '输入示例见exploratory_underwater__demo_inputs.csv，格式含case_id,Water,Cem。预测输出含MPa、缺失字段、超出训练单变量范围字段和探索状态。所有输入缺失仍可能有数值预测，但状态会明确标为all_inputs_missing，不应据此作科研判断。',
        '范围内不代表组合受训练支持或预测可靠；没有自动单位换算、概率置信区间或工程保证。模型仅加载本项目生成的本地joblib。','',
        '主要文件：results.json、training_oof_predictions.json、paired_imputation_comparison.json、cohort_and_training_folds.json、selected_models.json、*.joblib、*demo_inputs.csv、*demo_predictions.csv、verification.json、tests.log、training.log、console.log。',
        '新增代码与配置：configs/baseline_v1_1.json、src/concrete_audit/baseline_refinement.py、src/concrete_audit/predict.py、scripts/run_baseline_v1_1.py、scripts/predict_strength.py、scripts/summarize_baseline_v1_1.py、tests/test_baseline_refinement.py；README更新入口。',
        '函数说明：load_training只取锁定训练ID；ScaledKNNImputer按训练观察值定义距离尺度；pipeline封装插补和模型；run生成CV证据并拟合仅训练集模型；predict_csv白名单读入、范围标记及独立输出；summarize复核无留出样本参与。']
    (root/'report.md').write_text('\n'.join(lines)+'\n')
    print(json.dumps({k:v for k,v in verification.items() if k!='models'},ensure_ascii=False))


if __name__=='__main__':summarize()
