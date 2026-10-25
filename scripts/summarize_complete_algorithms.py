"""汇总五类算法、恒定列诊断和两篇PDF复查；重新校验全部评估产物。"""
from pathlib import Path
import json
import re
import numpy as np
from sklearn.impute import SimpleImputer
from sklearn.preprocessing import StandardScaler
from openpyxl import Workbook, load_workbook
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

from concrete_audit.baselines import regression_metrics
from concrete_audit.material_features import append_excel_row
from concrete_audit.nested_exploration import load_development, derive_ratios
from concrete_audit.parallel_exploration import digest, feature_matrix, write_json


def summarize():
    output=Path('outputs/complete-algorithm-comparison-v1');output.mkdir(exist_ok=False)
    paths={'trees':Path('outputs/parallel-nested-v2'),'svr_mlp':Path('outputs/parallel-svr-mlp-v1'),
           'filtered':Path('outputs/parallel-svr-mlp-variance-v1')}
    baseline=Path('outputs/parallel-exploration-v1');evidence=Path('outputs/two-paper-recheck-v1')
    results=[];predictions=[];traces=[];input_hashes={}
    for stage,path in paths.items():
        manifest=json.loads((path/'manifest.json').read_text())
        for p,expected in manifest['input_hashes_after'].items():
            if digest(p)!=expected:raise ValueError(f'输入变化：{p}')
        for filename in ['manifest.json','results.json','oof_predictions.json','inner_search_traces.json']:
            input_hashes[str(path/filename)]=digest(path/filename)
        results += [{**r,'stage':stage} for r in json.loads((path/'results.json').read_text())]
        predictions += json.loads((path/'oof_predictions.json').read_text())
        traces += json.loads((path/'inner_search_traces.json').read_text())
    prior_manifest,meta,rows,groups,y,folds=load_development(baseline)
    byid={r['sample_id']:(i,next(k for k,_,v in folds if i in v)) for i,r in enumerate(rows)}
    for r in results:
        pp=[p for p in predictions if p['experiment']==r['experiment']]
        assert len(pp)==len(rows) and {p['sample_id'] for p in pp}==set(byid)
        assert all(p['observed']==float(y[byid[p['sample_id']][0]]) and p['outer_fold']==byid[p['sample_id']][1] for p in pp)
        scores=regression_metrics([p['observed'] for p in pp],[p['predicted'] for p in pp])
        for key in ['MAE_MPa','RMSE_MPa','R2']:assert abs(scores[key]-r[key])<1e-9
    for trace in traces:
        t,v=set(trace['outer_train_groups']),set(trace['outer_validation_groups']);assert not t&v
        for inner in trace['inner_folds']:
            it,iv=set(inner['train_groups']),set(inner['validation_groups'])
            assert not it&iv and it|iv==t and not (it|iv)&v
    # 复现原始流水线的验证输入诊断，不使用验证目标拟合任何转换。
    config=json.loads(Path('configs/parallel_svr_mlp_v1.json').read_text())
    fields=prior_manifest['config']['feature_sets']['materials']
    matrices={'materials':(feature_matrix(rows,fields),fields),'materials_plus_ratios':derive_ratios(rows,fields,config)}
    diagnostics=[]
    for fs,(X,names) in matrices.items():
        for k,t,v in folds:
            imp=SimpleImputer(strategy='median',add_indicator=True,keep_empty_features=True)
            train=imp.fit_transform(X[t]);val=imp.transform(X[v]);scale=StandardScaler().fit(train)
            z=scale.transform(val);all_names=imp.get_feature_names_out(names)
            activated=[{'feature':str(all_names[j]),'training_constant_value':float(scale.mean_[j]),
                        'validation_max_abs_standardized':float(np.max(np.abs(z[:,j])))}
                       for j in range(len(all_names)) if scale.var_[j]==0 and np.max(np.abs(z[:,j]))>0]
            diagnostics.append({'feature_set':fs,'outer_fold':k,'constant_columns_activated_in_validation':activated,
                                'max_abs_validation_standardized':float(np.max(np.abs(z)))})
    write_json(output/'extrapolation_diagnostics.json',diagnostics)
    evidence_manifest=json.loads((evidence/'manifest.json').read_text())
    for p,expected in evidence_manifest['source_files_after_sha256'].items():assert digest(p)==expected
    composition=json.loads((evidence/'composition_recheck.json').read_text())
    test_log=Path('/tmp/underwater-complete-algorithms-tests.log').read_text()
    tests=re.search(r'Ran (\d+) tests in ([\d.]+)s\s+OK',test_log)
    if not tests:raise ValueError('全量测试尚未通过')
    write_json(output/'all_model_results.json',results)
    # 同材料字段的五算法主表；SVR/MLP使用已解释的折内恒定列过滤流水线。
    main=[]
    for model in ['rf','xgb','catboost','svr','mlp']:
        suffix='nested_filtered' if model in {'svr','mlp'} else 'nested'
        main.append(next(r for r in results if r['experiment']==f'{model}__materials__{suffix}'))
    write_json(output/'same_material_features_comparison.json',main)
    fig,ax=plt.subplots(figsize=(7,4));bars=ax.bar([r['model'] for r in main],[r['MAE_MPa'] for r in main])
    for bar,r in zip(bars,main):ax.text(bar.get_x()+bar.get_width()/2,bar.get_height()+.05,f'{r["MAE_MPa"]:.2f}',ha='center')
    ax.set_ylabel('Grouped OOF MAE (MPa)');ax.set_title('Same 283 samples and outer folds; model-specific pipelines')
    fig.tight_layout();fig.savefig(output/'five_models.png',dpi=160);plt.close(fig)
    wb=Workbook();wb.remove(wb.active)
    tables={'同材料特征五算法':(['model','experiment','n','MAE_MPa','RMSE_MPa','R2','macro_group_MAE_MPa'],main),
        '全部实验含失败结果':(['stage','experiment','n','MAE_MPa','RMSE_MPa','R2','macro_group_MAE_MPa','negative_predictions'],results),
        '逐条折外预测':(['experiment','source_row','paper','group','grade','outer_fold','observed','predicted','absolute_error'],predictions),
        '恒定列外推诊断':(['feature_set','outer_fold','max_abs_validation_standardized','constant_columns_activated_in_validation'],diagnostics),
        '两篇配比复查':(['source_row','paper','specimen_id','component_sum','printed_total','conflict','basis','water_value','water_unit','water_over_component_sum_candidate','notes'],composition),
        '内层搜索与参数':(['experiment','outer_fold','best_params','removed_constant_columns','outer_train_groups','outer_validation_groups'],traces)}
    for name,(keys,data) in tables.items():
        sheet=wb.create_sheet(name);append_excel_row(sheet,keys)
        for row in data:append_excel_row(sheet,[json.dumps(row.get(k),ensure_ascii=False) if isinstance(row.get(k),(list,dict)) else row.get(k) for k in keys])
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
    note=wb.create_sheet('解释与限制')
    for line in ['全部结果为探索性；同一批283条AB及外层折，目标/样本未变。',
        '过滤严格恒定列在内层训练折独立拟合，不根据验证值挑列。',
        '该对照在看到原始MLP异常后提出，仍存在跨轮适应性选择偏差。',
        '负预测和全部失败模型保留；未对预测作截断。',
        'Zhang相对水泥质量基准不能直接认证数据库绝对kg/m3。',
        'Al-Selmi分项1503与表注1502矛盾保留；不以算术替作者修正。',
        '145条历史留出本轮未参与拟合、转换、调参或评估；未产生最终独立测试结果。']:
        append_excel_row(note,[line])
    book=output/'五类模型对照与文献复查.xlsx';wb.save(book)
    check=load_workbook(book,read_only=True,data_only=False)
    assert not any(c.data_type=='f' for s in check for row in s for c in row);check.close()
    verify={'experiments_aggregated':len(results),'new_experiments':8,'new_oof_predictions':sum(p['experiment'].startswith(('svr','mlp')) for p in predictions),
        'sample_n':len(rows),'all_metrics_recalculated':True,'same_labels_and_outer_folds':True,
        'nested_source_isolation':True,'protected_inputs_unchanged':True,'unit_tests_passed':int(tests[1]),
        'pdf_records_rechecked':evidence_manifest['record_count'],'pdf_quote_matches':evidence_manifest['all_quotes_reproduced'],
        'pdf_conflict_records':sum(r['conflict'] for r in composition),'reserved_test_predictions':0,
        'input_hashes':input_hashes,'summary_script_sha256':digest(__file__),'adoption':'保留此前固定参数XGBoost作为探索参考，MLP不作为当前主模型'}
    write_json(output/'verification.json',verify)
    lines=['# 五类算法与两篇文献复查','',f'同一283条AB、20来源连通组、五外层折。新增SVR/MLP四组原始流水线和四组折内恒定列过滤对照；{tests[1]}项测试通过。',
        '输入和目标缩放均只在内层训练折拟合；MLP采用Dropout=.15、固定150轮，未用外层标签早停。','',
        '|模型（材料特征）|MAE MPa|RMSE MPa|R²|来源组平均MAE|','|---|---:|---:|---:|---:|']
    for r in main:lines.append(f'|{r["model"]}|{r["MAE_MPa"]:.3f}|{r["RMSE_MPa"]:.3f}|{r["R2"]:.3f}|{r["macro_group_MAE_MPa"]:.3f}|')
    lines+=['','异常与修正：',
        '- 原MLP材料流水线MAE52.114、R²-403.926；折内移除严格恒定列后MAE7.331、R²-0.493。原始失败结果完整保留。',
        '- 原因证据：外层折1的MK/SS在训练中无变化，验证标准化值最高分别480和914.4；MLP未学习的输入连接可能被激活。移除恒定列显著缓解此问题，但不证明所有错误只由该机制导致。',
        '- 剩余分布外输入、材料体系覆盖及样本来源不平衡仍存在。SVR加比值并过滤恒定列MAE6.392、RMSE10.671、R²0.144；不能只凭MAE认定所有指标都改善。',
        '- 这项诊断修正在看到旧外层结果后提出，不能将其称为预注册或独立最终验证。训练折过滤不使用验证信息，但研究过程仍有适应性选择偏差。',
        '- 本批模型均未超过此前固定参数XGBoost的MAE5.719；继续保留该模型作为探索参考，不宣称正式最优。','',
        '文献轨道：重新提取两篇原PDF指定页并复现两条不同配比引用，覆盖10条记录。Zhang按相对水泥质量归一化；Al-Selmi组分和1503与注释1502冲突仍保留。未修改原表、标签或正式可靠状态。','',
        '后续优先级：停止反复扩大同一开发集参数搜索；比较内层按来源组等权选参，补充同体系独立来源、完成优先字段核查后做数据版本同折对照。保留高强度记录，最终科研评价需另行预先确定独立验证方案。','',
        '修改/新增文件：configs/parallel_svr_mlp_v1.json、configs/parallel_svr_mlp_variance_v1.json、src/concrete_audit/neural_comparison.py、src/concrete_audit/scaled_comparison.py、scripts/run_svr_mlp.py、scripts/run_scaled_comparison.py、scripts/recheck_two_papers.py、scripts/summarize_complete_algorithms.py、tests/test_neural_comparison.py、tests/test_scaled_comparison.py、README.md、requirements-complete-models.txt。',
        '函数说明：search_model处理训练折内输入/目标缩放及搜索；filtered_search追加折内恒定列过滤；recheck从原PDF复现引用；summarize重算指标并关联预测、单位和来源证据。']
    (output/'report.md').write_text('\n'.join(lines)+'\n')
    print(json.dumps({k:v for k,v in verify.items() if k!='input_hashes'},ensure_ascii=False))


if __name__=='__main__':summarize()
