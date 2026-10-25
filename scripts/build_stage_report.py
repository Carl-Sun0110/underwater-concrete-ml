"""阶段材料：固定折外预测来源稳定性、教师汇报文档、实测可运行的CSV预测包。"""
from datetime import datetime, timezone
import csv
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml.ns import qn
from openpyxl import Workbook

from concrete_audit.parallel_exploration import digest, write_json
from concrete_audit.material_features import append_excel_row
from concrete_audit.stability import paired_source_stability


def build():
    root=Path('outputs/stage-report-v1');root.mkdir(exist_ok=False)
    baseline=Path('outputs/baseline_v1');refine=Path('outputs/baseline_v1_1');evidence=Path('outputs/priority-evidence-zero-v1')
    paths=[Path(__file__),Path('src/concrete_audit/stability.py'),baseline/'results.json',baseline/'sample_statistics.json',
        refine/'selected_models.json',refine/'training_oof_predictions.json',refine/'manifest.json',
        refine/'cohort_and_training_folds.json',evidence/'manifest.json',evidence/'paper_review_summary.json',
        Path('src/concrete_audit/predict.py'),Path('scripts/predict_strength.py')]
    previous=json.loads((refine/'manifest.json').read_text())
    for p,expected in previous['input_hashes_after'].items():
        if digest(p)!=expected:raise ValueError(f'历史输入发生变化：{p}')
    paths += [Path(p) for p in previous['input_hashes_after']]
    selected=json.loads((refine/'selected_models.json').read_text());paths += [Path(s['model_file']) for s in selected]
    before={str(p.resolve()):digest(p) for p in paths}
    predictions=json.loads((refine/'training_oof_predictions.json').read_text())
    summaries=[]
    for s in selected:
        model=[p for p in predictions if p['experiment']==s['experiment']]
        mean=[p for p in predictions if p['experiment']==s['cohort']+'__core__median__mean']
        stability=paired_source_stability(model,mean,resamples=2000,seed=42)
        summaries.append({'cohort':s['cohort'],'experiment':s['experiment'],**stability})
    write_json(root/'source_stability.json',summaries)
    expanded=next(s for s in summaries if s['cohort']=='exploratory_underwater')
    src=sorted(expanded['sources'],key=lambda s:s['delta_MAE_MPa'])
    fig,ax=plt.subplots(figsize=(9,7));deltas=[s['delta_MAE_MPa'] for s in src]
    ax.barh([s['group'] for s in src],deltas,color=['#467b9b' if d<0 else '#c57747' for d in deltas])
    ax.axvline(0,color='black',linewidth=.8);ax.set_xlabel('Model MAE - mean-reference MAE (MPa); lower is better')
    ax.set_title('Fixed training OOF predictions: benefit varies across source groups')
    fig.tight_layout();fig.savefig(root/'source_benefit.png',dpi=160);plt.close(fig)
    wb=Workbook();ws=wb.active;ws.title='来源稳定性'
    append_excel_row(ws,['cohort','group','n','papers','model_MAE_MPa','mean_reference_MAE_MPa','delta_MAE_MPa'])
    for s in summaries:
        for row in s['sources']:append_excel_row(ws,[s['cohort'],row['group'],row['n'],' | '.join(row['papers']),row['model_MAE_MPa'],row['mean_reference_MAE_MPa'],row['delta_MAE_MPa']])
    ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    ws=wb.create_sheet('重采样分位数');append_excel_row(ws,['cohort','metric','p2_5','median','p97_5','interpretation'])
    for s in summaries:
        for key,value in s['source_resampling_percentiles'].items():append_excel_row(ws,[s['cohort'],key,value['p2_5'],value['median'],value['p97_5'],s['interpretation']])
    wb.save(root/'来源稳定性诊断.xlsx')

    # 只打包已经发布的扩展候选，不发布未经PDF确认的零值假设模型。
    bundle=root/'prediction_bundle';bundle.mkdir();package=bundle/'src'/'concrete_audit';package.mkdir(parents=True)
    for p in Path('src/concrete_audit').glob('*.py'):shutil.copyfile(p,package/p.name)
    (bundle/'scripts').mkdir();shutil.copyfile('scripts/predict_strength.py',bundle/'scripts'/'predict_strength.py')
    chosen=next(s for s in selected if s['cohort']=='exploratory_underwater')
    shutil.copyfile(chosen['model_file'],bundle/'model.joblib')
    shutil.copyfile(refine/'exploratory_underwater__demo_inputs.csv',bundle/'example_inputs.csv')
    for filename in ['requirements.txt','requirements-modeling.txt','requirements-baseline-v1.txt']:
        shutil.copyfile(filename,bundle/filename)
    (bundle/'README.md').write_text('# 探索性配合比CSV预测包\n\n'
        '此包冻结Baseline V1.1扩展训练320条拟合的Water/Cem中位数XGBoost。它不是正式科研或工程验证模型。当前训练分组CV R²约0.084；未加入待核查Cem0假设。\n\n'
        '在本目录、Python 3.12环境运行：\n\n'
        '```bash\npython -m pip install -r requirements-baseline-v1.txt\nPYTHONPATH=src python scripts/predict_strength.py --model model.joblib --input example_inputs.csv --output your_predictions.csv\n```\n\n'
        '输入列case_id,Water,Cem；值须与原库名义配合比单位一致，不能自动把体积转换成质量。0/负数/空值依旧标为未知。结果包含缺失、超训练单变量范围和探索状态；范围内也不能证明可靠，不提供置信区间。\n\n'
        'example_inputs.csv来自现有训练样本，只是程序演示；predictions_from_package_test.csv是打包后实际运行结果，不能当独立测试。输出文件必须不存在。\n')
    env=os.environ.copy();env['PYTHONPATH']=str((bundle/'src').resolve())
    command=[sys.executable,'scripts/predict_strength.py','--model','model.joblib','--input','example_inputs.csv','--output','predictions_from_package_test.csv']
    completed=subprocess.run(command,cwd=bundle.resolve(),env=env,text=True,capture_output=True,timeout=60,check=True)
    (bundle/'package_test.log').write_text(completed.stdout+completed.stderr)
    with (bundle/'predictions_from_package_test.csv').open(encoding='utf-8-sig') as f:packaged=list(csv.DictReader(f))
    with (refine/'exploratory_underwater__demo_predictions.csv').open(encoding='utf-8-sig') as f:original=list(csv.DictReader(f))
    assert packaged==original and len(packaged)==3
    assert digest(bundle/'model.joblib')==digest(chosen['model_file'])
    write_json(bundle/'package_manifest.json',{'source_model_sha256':digest(chosen['model_file']),
        'source_experiment':chosen['experiment'],'package_inference_test':'passed','predictions_equal_original_demo':True,
        'demo_rows_are_training_replays':True,'files':{str(p.relative_to(bundle)):digest(p) for p in bundle.rglob('*') if p.is_file() and '__pycache__' not in p.parts}})
    with zipfile.ZipFile(root/'prediction_bundle.zip','x',compression=zipfile.ZIP_DEFLATED) as archive:
        for p in bundle.rglob('*'):
            if p.is_file() and '__pycache__' not in p.parts:archive.write(p,p.relative_to(root))
    with zipfile.ZipFile(root/'prediction_bundle.zip') as archive:
        packaged_manifest=json.loads(archive.read('prediction_bundle/package_manifest.json'))
        assert set(archive.namelist())=={'prediction_bundle/'+p for p in packaged_manifest['files']}|{'prediction_bundle/package_manifest.json'}
        for name,expected in packaged_manifest['files'].items():
            assert hashlib.sha256(archive.read('prediction_bundle/'+name)).hexdigest()==expected
        with tempfile.TemporaryDirectory(prefix='underwater-package-test-') as folder:
            archive.extractall(folder);extracted=Path(folder)/'prediction_bundle'
            isolated_env=os.environ.copy();isolated_env['PYTHONPATH']=str(extracted/'src');isolated_env['PYTHONDONTWRITEBYTECODE']='1'
            replay=subprocess.run([sys.executable,'scripts/predict_strength.py','--model','model.joblib',
                '--input','example_inputs.csv','--output','predictions_extracted_test.csv'],cwd=extracted,
                env=isolated_env,text=True,capture_output=True,check=True,timeout=60)
            with (extracted/'predictions_extracted_test.csv').open(encoding='utf-8-sig') as f:assert list(csv.DictReader(f))==original
            (root/'zip_extract_test.log').write_text(replay.stdout+replay.stderr)

    statistics=json.loads((baseline/'sample_statistics.json').read_text());metrics=json.loads((baseline/'results.json').read_text())
    delta=expanded['source_resampling_percentiles']['paired_macro_MAE_delta_MPa']
    interval=expanded['source_resampling_percentiles']['model_sample_MAE_MPa']
    blocks=[]
    def paragraph(text):blocks.append(('paragraph',text))
    def heading(text):blocks.append(('heading',text))
    def table(headers,rows):blocks.append(('table',(headers,rows)))
    heading('水下混凝土28天抗压强度预测：阶段性探索建模报告')
    paragraph('日期：2026年10月24日。目的：说明已运行的建模流程、真实验证表现、来源稳定性和目前证据边界。本文不作为最终科研结论。')
    heading('1. 当前进展与可支持的结论')
    paragraph('已跑通只读导入、筛选与版本记录、缺失输入处理、RF/XGBoost/CatBoost训练、随机配合比簇与论文分组评价、残差和来源诊断、CSV批量预测。当前跨文献性能不稳定，尚不能支持可靠高精度泛化的结论。')
    paragraph('原始Excel和强度标签未覆盖。数据及部分来源已经过多轮探索，不能把现有开发分数称作全新盲测。没有新增正式“已核实可用”样本。')
    heading('2. 实际文件及数据范围')
    paragraph('未找到《水下混凝土_人工核查版.xlsx》和《水下混凝土_筛选结果_20261012.xlsx》。实际读取原版UnderwaterConcrete - Total Data (Select 0927).xlsx；原版Selected、批注核对版及筛选方案是不同文件，不冒充缺失版本。')
    table(['项目','实际记录数'],[['原版Total',statistics['source_records']],['原表标为水下',statistics['underwater_records']],
        ['水下Converted正值',statistics['positive_converted_underwater']],['有直接水下成型辅助证据的主集',statistics['primary_direct_evidence']],
        ['扩展探索集（包含主集）',statistics['exploratory_total']],['额外范围待核查',statistics['unknown_scope_additional']]])
    paragraph('主集成型范围有原文辅助核查支持，不等于整条数据或换算依据已正式人工核实。使用原表存储Converted CS_28d；0.95/0.8/0.76系数沿用用户许可的暂定规则。算术一致不证明可靠。已知非直接成型的30条单独隔离，未补造目标。')
    heading('3. Baseline V1方法与真实结果')
    paragraph('原材料11字段仅取配合比参数；Density、其他实测性能和目标派生变量未纳入。零值含义未明按未知，中位数与缺失指示器只在训练折拟合。固定参数、无搜索；划分前识别近重复配比，不按目标删样本。')
    comparison=[r for r in metrics if r['cohort']=='primary_direct' and r['model']!='mean']
    table(['主集方案','模型','评价数','MAE/MPa','RMSE/MPa','R²'],[[r['protocol'],r['model'],r['n'],f'{r["MAE_MPa"]:.3f}',f'{r["RMSE_MPa"]:.3f}',f'{r["R2"]:.3f}'] for r in comparison])
    paragraph('随机划分按近重复配比簇隔离，但允许同文献不同配比跨集。论文留出按来源连通组隔离；两种划分的测试分布不同，差距不能全部归因于泄漏或过拟合。论文留出R²均为负，提示跨来源迁移困难。主集只有6个来源连通组，五折验证样本数也不均。')
    heading('4. Baseline V1.1：仅训练集内部选择')
    paragraph('锁定V1论文留出。本轮原有训练内比较只用主集117条/4来源组、扩展训练320条/23来源组；95/147条留出不再参与选方案或新评价。比较Water/Cem两字段与11字段、中位数与标准化KNN；按来源组平均MAE选候选。')
    table(['训练队列','候选','MAE/MPa','RMSE/MPa','R²'],[[s['cohort'],s['model']+' / '+s['feature_set']+' / '+s['imputation'],f'{s["MAE_MPa"]:.3f}',f'{s["RMSE_MPa"]:.3f}',f'{s["R2"]:.3f}'] for s in selected])
    paragraph('主集训练只余4组，其中Nasr组104条，留该组验证时只有其他体系13条可训练，结果极不稳定。这些训练CV分数不能与V1全体五折直接比较。模型选择与报告复用了同一训练CV，仍可能乐观。')
    heading('5. 本轮新增：固定折外预测的来源稳定性')
    paragraph(f'扩展候选在23个来源连通组中，{expanded["sources_improved"]}组MAE优于均值参照，{expanded["sources_worsened"]}组更差。对来源组进行2000次配对重采样、每次保留组内所有记录，既不重训练，也不挑测试集。')
    paragraph(f'模型样本平均MAE为{expanded["model_sample_MAE_MPa"]:.3f} MPa，来源重采样2.5%–97.5%分位范围为{interval["p2_5"]:.3f}–{interval["p97_5"]:.3f} MPa。来源宏平均MAE相对均值参照的差值范围为{delta["p2_5"]:.3f}至{delta["p97_5"]:.3f} MPa，跨过0，说明收益对来源构成敏感。')
    paragraph('上述范围是固定训练折外误差的描述性稳定性诊断，不是独立测试置信区间，不是p值或泛化保证。它不包含重新拟合模型的全部不确定性，也不能消除多轮探索和选模偏差。')
    blocks.append(('image',root/'source_benefit.png'))
    heading('6. 文献证据缺口与数据语义')
    paragraph('已整理4篇缺原文来源的26条记录。当前网络代理请求返回403，没有取得新PDF；题录、DOI、备注和单元格内容是原表线索，未冒充原文验证。')
    paragraph('Zhongzhe Zhang 2023的747/751/755/759行Cem0、GGBS550、AlkaliAct189，可能为真实无水泥体系；需确认其水泥、激发剂溶液及有效水含义。751行Water64与其他相关行128也保留差异，不按邻行自动改写。')
    paragraph('Wang 2024的790行涉及砂浆灌入预置骨料，强度标签对象待确认；Horszczaruk 2017原表备注明示水下成型为推定；Horszczaruk 2023的压力养护/成型条件仍待核。高误差是复核线索，不直接等同错误数据。')
    paragraph('只对上述4条Cem0做假设敏感性时，两特征XGBoost训练CV MAE9.118→9.018 MPa。这个小变化不能确认真零；原始数据、标签、正式状态和已发布预测模型未被假设替换。')
    heading('7. 下一阶段的具体交付')
    paragraph('先取得Zhongzhe Zhang等2023年Effect of antiwashout admixture on compressive strength of alkali-activated underwater non-dispersible concrete。原表DOI 10.1680/jmacr.22.00239，本轮未联网核验。优先比对A/C系列配合比表、激发剂溶液及水量定义、原始28天强度和150mm试件依据。无需先收齐全部论文。')
    paragraph('核实后生成新的输入版本并做固定训练折对照；现有标签保留，改动须有原文位置与换算依据。同时补充独立材料体系来源，预先确定后续独立验证方案，持续维持文献核查和探索建模两轨。')
    heading('8. 可复现程序与运行边界')
    paragraph('prediction_bundle.zip包含冻结的两特征XGBoost、CSV读入代码、依赖与示例。打包后在独立子进程实际执行成功，输出与原演示逐项一致。示例来自训练记录，不是新实验，也不是新准确率证据。')
    paragraph('预测会输出缺失输入、超训练单变量范围和探索状态；范围内不证明组合受训练支持。单位不自动换算，未输出置信区间，尚不具备工程配合比设计验证条件。')
    paragraph('主要证据：Baseline V1与V1.1结果目录、来源稳定性诊断.xlsx、source_stability.json、优先文献核查包，以及所有输入哈希、配置和测试日志。')
    markdown=[];doc=Document()
    style=doc.styles['Normal'];style.font.name='SimSun';style.font.size=Pt(10.5);style._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
    for kind,value in blocks:
        if kind=='heading':markdown.extend(['## '+value,'']);doc.add_heading(value,level=1)
        elif kind=='paragraph':markdown.extend([value,'']);doc.add_paragraph(value)
        elif kind=='table':
            headers,rows=value;markdown.extend(['|'+ '|'.join(map(str,headers))+'|','|'+'|'.join(['---']*len(headers))+'|'])
            table_doc=doc.add_table(rows=1,cols=len(headers));table_doc.style='Light Shading Accent 1'
            for j,h in enumerate(headers):table_doc.rows[0].cells[j].text=str(h)
            for row in rows:
                markdown.append('|'+ '|'.join(map(str,row))+'|');cells=table_doc.add_row().cells
                for j,v in enumerate(row):cells[j].text=str(v)
            markdown.append('')
        else:markdown.extend(['![来源组收益](source_benefit.png)','']);doc.add_picture(str(value),width=Inches(6))
    (root/'阶段性探索建模报告.md').write_text('\n'.join(markdown)+'\n');doc.save(root/'阶段性探索建模报告.docx')
    (root/'给老师的简短汇报.txt').write_text('老师您好，目前已跑通水下混凝土28天抗压强度预测流程，完成随机森林、XGBoost、CatBoost基线及按论文分组评价。随机划分表现较好，但跨文献泛化仍不稳定。现有扩展训练候选在23个来源组中12组优于均值参照、11组更差，因此暂不将结果作为可靠高精度科研结论。下一步拟优先核实Zhongzhe Zhang 2023的无水泥候选配比、激发剂水量及试件换算依据，并补充独立来源。已附阶段报告和可运行程序，原始数据及标签未覆盖，待核查记录均保留标记。\n')
    testlog=Path('/tmp/underwater-stability-report-tests.log').read_text();tests=re.search(r'Ran (\d+) tests in ([\d.]+)s\s+OK',testlog)
    if not tests:raise ValueError('全量测试未通过')
    shutil.copyfile('/tmp/underwater-stability-report-tests.log',root/'tests.log')
    after={p:digest(p) for p in before}
    if after!=before:raise RuntimeError('旧数据、模型或证据发生变化')
    write_json(root/'manifest.json',{'created_utc':datetime.now(timezone.utc).isoformat(),'input_hashes_before':before,'input_hashes_after':after,
        'new_model_fits':0,'new_heldout_predictions':0,'source_resamples_per_cohort':2000,'unit_tests_passed':int(tests[1]),
        'package_test_passed':True,'zip_file_hashes_verified':True,'extracted_package_test_passed':True,
        'package_demo_prediction_rows':3,'package_demo_is_training_replay':True,
        'new_pdf_evidence':0,'formal_status_changes':0,'report_docx_sha256':digest(root/'阶段性探索建模报告.docx'),
        'bundle_zip_sha256':digest(root/'prediction_bundle.zip')})
    print(json.dumps({'source_groups':expanded['source_groups'],'improved':expanded['sources_improved'],
        'worsened':expanded['sources_worsened'],'paired_macro_delta_percentiles':delta,'tests_passed':int(tests[1]),'package_test':'passed'},ensure_ascii=False))


if __name__=='__main__':build()
