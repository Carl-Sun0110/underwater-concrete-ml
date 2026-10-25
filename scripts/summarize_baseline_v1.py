"""生成Baseline V1真实预测/残差图、来源诊断、Excel及实施总结。"""
import argparse
from collections import Counter, defaultdict
import json
from pathlib import Path
import re
import shutil

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import numpy as np
from openpyxl import Workbook, load_workbook

from concrete_audit.baseline_v1 import read_source
from concrete_audit.baselines import regression_metrics
from concrete_audit.material_features import append_excel_row
from concrete_audit.parallel_exploration import digest, numeric, write_json


def summarize(root,test_log,console_log):
    root=Path(root);plots=root/'plots';plots.mkdir(exist_ok=False)
    manifest=json.loads((root/'manifest.json').read_text());config=manifest['config']
    results=json.loads((root/'results.json').read_text());pred=json.loads((root/'predictions.json').read_text())
    splits=json.loads((root/'split_manifest.json').read_text());stats=json.loads((root/'sample_statistics.json').read_text())
    selection=json.loads((root/'selection_records.json').read_text());selected=json.loads((root/'modeling_records.json').read_text())
    lookup={r['sample_id']:r for r in selected};features=json.loads((root/'features_and_missingness.json').read_text())
    imps=json.loads((root/'permutation_importance.json').read_text());duplicates=json.loads((root/'duplicate_analysis.json').read_text())
    for p,expected in manifest['input_hashes_after'].items():assert digest(p)==expected
    for split in splits:
        assert not set(split['train_ids'])&set(split['test_ids'])
        assert not set(split['train_recipe_groups'])&set(split['test_recipe_groups'])
        if split['protocol'].startswith('paper'):assert not set(split['train_paper_groups'])&set(split['test_paper_groups'])
    for result in results:
        pp=[p for p in pred if (p['cohort'],p['protocol'],p['model'])==(result['cohort'],result['protocol'],result['model'])]
        assert len(pp)==result['n']==len({p['sample_id'] for p in pp})
        assert all(p['actual']==lookup[p['sample_id']]['target'] for p in pp)
        score=regression_metrics([p['actual'] for p in pp],[p['predicted'] for p in pp])
        for k in ['MAE_MPa','RMSE_MPa','R2']:assert abs(result[k]-score[k])<1e-9
        for p in pp:
            split=next(s for s in splits if (s['cohort'],s['protocol'],s['fold'])==(p['cohort'],p['protocol'],p['fold']))
            assert p['sample_id'] in split['test_ids']
    # 重读源表验证建模值，不以缓存快照冒充Excel检查；记录所有字段实际类型。
    source,dictionary=read_source(config['raw']);source_byid={r['sample_id']:r for r in source}
    for r in selected:assert r['target']==numeric(source_byid[r['sample_id']]['original_values'][config['target']])
    field_types=[]
    for field in dictionary:
        values=[r['original_values'].get(field) for r in source];counts=Counter()
        for v in values:
            if v is None or v in ('','/','-'):counts['unknown_marker']+=1
            elif isinstance(v,str) and v.startswith('='):counts['excel_formula_no_recalculation']+=1
            elif numeric(v) is not None:counts['numeric']+=1
            else:counts['text_or_category']+=1
        field_types.append({'field':field,'declared_unit':dictionary[field].get('unit'),'observed_types':dict(counts)})
    write_json(root/'field_types.json',field_types)
    availability=[]
    for cohort in ['primary_direct','exploratory_underwater']:
        data=[r for r in selected if cohort=='exploratory_underwater' or r['direct_scope_evidence']]
        for field in config['features']:
            n=sum(r['features'][field] is not None for r in data)
            availability.append({'cohort':cohort,'field':field,'n':len(data),'observed_positive_n':n,'unknown_n':len(data)-n,'unknown_fraction':1-n/len(data)})
    write_json(root/'feature_availability_by_cohort.json',availability)
    models=[m for m in config['models'] if any(r['model']==m for r in results)]
    for cohort in ['primary_direct','exploratory_underwater']:
        for protocol in ['random_recipe_holdout','paper_group_holdout','paper_group_cv']:
            fig,axes=plt.subplots(1,len(models),figsize=(5*len(models),4.5),squeeze=False)
            fig2,axes2=plt.subplots(2,len(models),figsize=(5*len(models),7),squeeze=False)
            for j,model in enumerate(models):
                pp=[p for p in pred if (p['cohort'],p['protocol'],p['model'])==(cohort,protocol,model)]
                if not pp:continue
                actual=np.array([p['actual'] for p in pp]);prediction=np.array([p['predicted'] for p in pp]);residual=actual-prediction
                metric=next(r for r in results if (r['cohort'],r['protocol'],r['model'])==(cohort,protocol,model))
                ax=axes[0,j];ax.scatter(actual,prediction,s=16,alpha=.65)
                low=min(actual.min(),prediction.min());high=max(actual.max(),prediction.max());ax.plot([low,high],[low,high],'k--')
                ax.set_xlabel('Actual stored Converted CS_28d (MPa)');ax.set_ylabel('Predicted (MPa)')
                ax.set_title(f'{model}: n={len(pp)}, MAE={metric["MAE_MPa"]:.2f}, R2={metric["R2"]:.3f}')
                axes2[0,j].scatter(prediction,residual,s=16,alpha=.65);axes2[0,j].axhline(0,color='black',linestyle='--')
                axes2[0,j].set_xlabel('Predicted (MPa)');axes2[0,j].set_ylabel('Actual - predicted (MPa)');axes2[0,j].set_title(model)
                axes2[1,j].hist(residual,bins=min(20,max(5,len(pp)//5)));axes2[1,j].set_xlabel('Residual (MPa)');axes2[1,j].set_ylabel('Count')
            fig.suptitle(f'{cohort}: {protocol}');fig.tight_layout();fig.savefig(plots/f'{cohort}__{protocol}__actual_vs_predicted.png',dpi=150);plt.close(fig)
            fig2.suptitle(f'{cohort}: {protocol}');fig2.tight_layout();fig2.savefig(plots/f'{cohort}__{protocol}__residuals.png',dpi=150);plt.close(fig2)
    paper_groups=defaultdict(list)
    for p in pred:
        if p['protocol']=='paper_group_cv' and p['model']!='mean':paper_groups[p['cohort'],p['model'],p['paper']].append(p)
    paper_errors=[]
    for (cohort,model,paper),pp in paper_groups.items():
        paper_errors.append({'cohort':cohort,'model':model,'paper':paper,'n':len(pp),
            **regression_metrics([p['actual'] for p in pp],[p['predicted'] for p in pp]),
            'bias_pred_minus_actual_MPa':float(np.mean([p['predicted']-p['actual'] for p in pp])),
            'source_rows':[p['source_row'] for p in pp]})
    paper_errors.sort(key=lambda r:r['MAE_MPa'],reverse=True);write_json(root/'paper_error_analysis.json',paper_errors)
    bysample=defaultdict(list)
    for p in pred:
        if p['cohort']=='exploratory_underwater' and p['protocol']=='paper_group_cv' and p['model']!='mean':bysample[p['sample_id']].append(p)
    high_error=[]
    for sample_id,pp in bysample.items():
        r=lookup[sample_id];raw=source_byid[sample_id]['original_values']
        high_error.append({'sample_id':sample_id,'source_row':r['source_row'],'paper':r['paper'],'target':r['target'],
            'mean_absolute_error_across_models_MPa':float(np.mean([abs(p['actual']-p['predicted']) for p in pp])),
            'model_predictions':{p['model']:p['predicted'] for p in pp},'label_extraction':raw.get('DataSource'),
            'raw_CS_28d':raw.get('CS_28d'),'factor':raw.get('Conversion Factor'),'specimen_geometry':raw.get('CS_Spec'),
            'unknown_features':[f for f,v in r['features'].items() if v is None],
            'action':'核查线索，非已证实异常；不按模型误差更改标签或删除样本'})
    high_error.sort(key=lambda r:r['mean_absolute_error_across_models_MPa'],reverse=True)
    write_json(root/'high_error_records.json',high_error[:30])
    priority=[]
    for field in config['features']:
        a=next(r for r in availability if r['cohort']=='exploratory_underwater' and r['field']==field)
        scores=[r['MAE_increase_by_feature'][field] for r in imps if r['cohort']=='exploratory_underwater' and r['protocol']=='paper_group_holdout']
        priority.append({'field':field,'unknown_fraction':a['unknown_fraction'],'mean_holdout_permutation_MAE_increase_MPa':float(np.mean(scores)) if scores else None,
            'caveat':'描述性测试集置换，相关特征可互相替代；负重要性或稀疏性不自动触发删除'})
    write_json(root/'feature_review_priorities.json',priority)
    log=Path(test_log).read_text();tests=re.search(r'Ran (\d+) tests in ([\d.]+)s\s+OK',log)
    if not tests:raise ValueError('未找到本次全部测试通过结果')
    shutil.copyfile(test_log,root/'tests.log');shutil.copyfile(console_log,root/'console.log')
    for file in ['scripts/run_baseline_v1.py',__file__]:shutil.copyfile(file,root/'code'/Path(file).name)
    wb=Workbook();wb.remove(wb.active)
    tables={'指标':(['cohort','protocol','model','cohort_n','n','MAE_MPa','RMSE_MPa','R2','macro_source_MAE_MPa','mean_training_MAE_MPa','mean_training_R2'],results),
        '建模数据':(['sample_id','source_row','paper','target','direct_scope_evidence','label_status','features'],selected),
        '全部筛选记录':(['source_row','paper','specimen_id','casting_environment','raw_CS_28d','stored_Converted_CS_28d','factor','scope_state','included_primary','included_exploratory','reasons','label_status'],selection),
        '特征与缺失':(['cohort','field','n','observed_positive_n','unknown_n','unknown_fraction'],availability),
        '真实与预测':(['cohort','protocol','model','fold','source_row','paper','actual','predicted','residual_actual_minus_predicted'],pred),
        '论文误差':(['cohort','model','paper','n','MAE_MPa','RMSE_MPa','R2','bias_pred_minus_actual_MPa','source_rows'],paper_errors),
        '高误差待核查':(['source_row','paper','target','mean_absolute_error_across_models_MPa','raw_CS_28d','factor','specimen_geometry','label_extraction','unknown_features','action'],high_error[:30]),
        '特征核查优先级':(['field','unknown_fraction','mean_holdout_permutation_MAE_increase_MPa','caveat'],priority)}
    for title,(keys,data) in tables.items():
        ws=wb.create_sheet(title);append_excel_row(ws,keys)
        for row in data:append_excel_row(ws,[json.dumps(row.get(k),ensure_ascii=False) if isinstance(row.get(k),(list,dict)) else row.get(k) for k in keys])
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    wb.save(root/'Baseline_V1结果与筛选记录.xlsx')
    check=load_workbook(root/'Baseline_V1结果与筛选记录.xlsx',read_only=True,data_only=False)
    assert not any(c.data_type=='f' for s in check for row in s for c in row);check.close()
    verification={'model_metric_rows':len(results),'prediction_rows':len(pred),'unit_tests_passed':int(tests[1]),
        'raw_workbook_rows_rechecked':len(source),'all_targets_match_stored_converted':True,'all_metrics_recomputed':True,
        'recipe_groups_isolated':True,'paper_groups_isolated_in_grouped_protocols':True,'protected_inputs_unchanged':True,
        'actual_prediction_plots':6,'residual_plots':6,'summary_script_sha256':digest(__file__)}
    write_json(root/'verification.json',verification)
    lines=['# Baseline V1 第一轮实际建模总结（2026-10-24）','',
        '本报告为本次明确协议的独立Baseline V1运行，不沿用历史分数。数据及部分来源此前已参加探索，不能称全新盲测；本轮测试评价后不再用于反复选参并声称独立。','',
        '## 文件与实际数据','',
        '未找到《水下混凝土_人工核查版.xlsx》《水下混凝土_筛选结果_20261012.xlsx》。实际读取原版Total；Selected为编码版，批注版与筛选方案为独立版本，不冒充所缺文件。工作簿结构、原始附件哈希和所有项目Excel清单见excel_inventory.json。',
        f'Total {stats["source_records"]}条，水下标记{stats["underwater_records"]}条；其中原始28天强度正值{stats["positive_raw_28d_underwater"]}条，Converted正值{stats["positive_converted_underwater"]}条。',
        f'497条满足本轮暂定换算/一致性候选，其中30条已知非直接成型单独隔离。最终主集{stats["primary_direct_evidence"]}条、扩展探索集{stats["exploratory_total"]}条（包含主集及额外范围待核查{stats["unknown_scope_additional"]}条）。',
        '原文辅助证据支持主集成型范围，整条标签/换算仍非正式人工核实。使用原表存储Converted，不重算标签、不补目标；0.95/0.8/0.76为用户既有暂定许可，不因算术一致就认定可靠。未强行纳入其他系数对应记录，也未判为伪造。','',
        '## 特征、划分与固定参数','',
        '11项输入：'+', '.join(config['features'])+'。材料正值保留，0含义未明按未知；中位数插补加缺失指示器仅训练折拟合，整列未知的0是数值占位。未使用Density、强度衍生变量、其他实测性能；SP/AWA/纤维及w/c、w/b的单位、掺量基准或公式来源待统一，未在本轮直接混用。',
        '三模型固定150棵/轮；XGBoost/CatBoost深度4、学习率.04，RF叶节点最少3；无参数搜索。增加均值参照，完整参数见config.json、代码与模型文件。',
        '随机种子42。随机对照先按近重复配比簇划分，允许同论文不同配比跨集；未做有近重复泄漏的逐行随机。组别验证按论文/研究组及跨论文近重复连通组隔离。近重复为全配比字段已知模式相同且逐项1%以内，不看目标，不据此删除记录。','',
        '|数据集|方案|训练数|验证/测试数|','|---|---|---:|---:|']
    for s in splits:
        if s['protocol'].endswith('holdout'):lines.append(f'|{s["cohort"]}|{s["protocol"]}|{s["train_n"]}|{s["test_n"]}|')
    lines+=['','## 真实评价','', '|数据集|方案|模型|n|MAE MPa|RMSE MPa|R²|','|---|---|---|---:|---:|---:|---:|']
    for r in results:lines.append(f'|{r["cohort"]}|{r["protocol"]}|{r["model"]}|{r["n"]}|{r["MAE_MPa"]:.3f}|{r["RMSE_MPa"]:.3f}|{r["R2"]:.3f}|')
    lines+=['','## 解释、限制与下一轮','',
        '- 主集随机配比簇划分表现明显好于论文留出。随机对照虽隔离了近重复配比，仍可学习同论文体系特征；按论文留出存在材料体系、强度范围及实验条件迁移。差距不能全部归因于泄漏或训练过拟合，也不能把随机R²当作跨文献能力。',
        '- 主集仅6个来源连通组、7个论文文件名，五折中验证数不均；扩展集29组。组测试比例按组抽取，样本比例不一定20%，不得为改善指标重新挑测试论文。各折、训练分数、来源宏平均误差均保留。',
        '- 论文分组固定留出三模型R²均为负，部分模型不及均值参照；样本强度方差较小也影响R²。主集五折的三指标排名不一致，暂不宣布最终最优模型。',
        '- 来源覆盖、单位、密度假设和未知零值语义仍待核；Converted使用暂定系数，需按原试件尺寸、形状、龄期、养护条件和换算文献逐条复核，不用预测值倒推标签。',
        '- 高误差记录及论文按分组折外预测产生；高强度值只是复核线索，不自动认定错误。优先核对高影响来源的配合比绝对/相对基准、砂/粉煤灰字段和标签图形读取。',
        '- 特征严重缺失与重要性见feature_review_priorities.json。测试置换仅用于本轮描述，相关列可相互替代，不能根据这次测试重要性删列后继续称独立测试。',
        '- 下一轮：固定本轮协议，在训练来源内部比较KNN/中位数及单位一致的配比比值；扩充独立来源、进行标签换算敏感性、按预先确定方案开展新验证。SHAP本轮未强制追加。',
        f'- {tests[1]}项全量测试通过；原始输入哈希未变，全部指标重算验证通过；12张预测/残差图已生成。','',
        '## 文件与复现','',
        '结果目录含Excel/metrics.csv、selection_records.json、modeling_records.json、features_and_missingness.json、preprocessing.json、split_manifest.json、duplicate_analysis.json、predictions.json、paper_error_analysis.json、high_error_records.json、plots/、模型joblib、code/、training.log、console.log、tests.log和manifest.json。',
        '仓库新增configs/baseline_v1.json、src/concrete_audit/baseline_v1.py、scripts/run_baseline_v1.py、scripts/summarize_baseline_v1.py、tests/test_baseline_v1.py；README更新本轮入口。',
        '```bash','PYTHONPATH=src /workspace/.venv-underwater-ml/bin/python scripts/run_baseline_v1.py --output outputs/baseline_v1_rerun',
        '```','',
        '输出目录必须不存在。关键函数：read_source映射参数字典及身份合并；select_records保留逐条理由和存储目标；clusters在划分前识别近重复；model_pipeline封装训练折插补；run执行固定参数评估；summarize重算指标并导出图表。']
    (root/'report.md').write_text('\n'.join(lines)+'\n')
    (root/'README.md').write_text('本次Baseline V1结果见report.md；Excel为Baseline_V1结果与筛选记录.xlsx；图在plots/，代码副本在code/。\n全部结果为探索性，原始数据和历史输出未覆盖。\n')
    print(json.dumps(verification,ensure_ascii=False))


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--run',default='outputs/baseline_v1');p.add_argument('--test-log',default='/tmp/underwater-baseline-v1-tests.log')
    p.add_argument('--console-log',default='/tmp/underwater-baseline-v1-console.log');a=p.parse_args();summarize(a.run,a.test_log,a.console_log)
