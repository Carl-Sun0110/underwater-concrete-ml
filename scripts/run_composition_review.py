"""实跑组成、单位和来源混杂诊断；只写新报告，不改训练或原始数据。"""
from pathlib import Path
from collections import Counter
import argparse
import json
import subprocess
import sys
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'src'))
from concrete_audit.composition_review import derive_composition, family_confounding
from concrete_audit.field_review import digest, validate_evidence
from concrete_audit.evidence_features import fold_support
from concrete_audit.material_features import append_excel_row


def run(output):
    output = Path(output)
    config_path = ROOT / 'configs/composition_review_evidence.json'
    config = json.loads(config_path.read_text())
    scope = ROOT / 'outputs/casting-scope-gate-v1'
    ids = json.loads((scope / 'scope_views.json').read_text())['strict_evidence_scope']['train_sample_ids']
    gate = json.loads((scope / 'scope_registry.json').read_text())
    folds = json.loads((scope / 'training_plan.json').read_text())['plans']['strict_evidence_scope']['cv']
    split_path = ROOT / 'outputs/exploratory-baselines-v2/split_manifest.json'
    split = json.loads(split_path.read_text())['rows']
    groups = {r['sample_id']: r['leakage_group'] for r in split}
    test_ids = {r['sample_id'] for r in split if r['split'] == 'test'}
    labels_path = ROOT / 'outputs/provisional-labels-v1/exploratory_labels.json'
    labels = [r for r in json.loads(labels_path.read_text()) if r['sample_id'] in set(ids)]
    by_row = {r['source_row']: r for r in labels}
    if len(by_row) != len(ids) or set(ids) & test_ids:
        raise ValueError('组成核查身份不匹配')
    protected = [config_path, labels_path, split_path, scope / 'scope_views.json', scope / 'scope_registry.json',
                 scope / 'training_plan.json', *map(Path, config['protected_inputs']),
                 *[Path(p['path']) for p in config['papers'].values()]]
    before = dict(gate['protected_inputs_sha256'])
    if any(digest(p) != h for p, h in before.items()):
        raise ValueError('上游数据版本已改变')
    before.update({str(p.resolve()): digest(p) for p in protected})
    pages = {}
    for paper, metadata in config['papers'].items():
        if digest(metadata['path']) != metadata['sha256']:
            raise ValueError('PDF版本改变')
        pages[paper] = subprocess.run(['pdftotext', '-layout', metadata['path'], '-'], capture_output=True, check=True).stdout.decode().split('\f')
    records, seen = [], set()
    for recipe in config['annotations']:
        row = recipe['source_row']
        if row in seen or row not in by_row:
            raise ValueError('组成核查存在重复或非训练行')
        seen.add(row)
        src = by_row[row]
        if (recipe['specimen_id'], recipe['registry_pdf_filename']) != (src['original_values']['Specimen ID'], src['original_values']['PDF filename']):
            raise ValueError('组成证据身份不一致')
        validate_evidence(recipe, pages[recipe['paper']], config['papers'][recipe['paper']]['sha256'])
        record = {'sample_id': src['sample_id'], 'source_row': row, 'paper': recipe['paper'],
                  'declared_mix_family': recipe['declared_mix_family'], 'components': recipe['components'],
                  'fibers': recipe['fibers'], 'source_notes': recipe['notes'],
                  'prior_reliability_status': src['prior_reliability_status'], 'human_confirmation': '待核查',
                  **derive_composition(recipe)}
        records.append(record)
    if set(seen) != set(by_row):
        raise ValueError('组成核查未覆盖固定训练记录')
    confounding = family_confounding(records, groups, folds)
    candidates = [{'sample_id': r['sample_id'], 'candidate_values': r['candidate_ratios']} for r in records]
    coverage = fold_support(candidates, ids, groups, folds, list(records[0]['candidate_ratios']))
    summary = {'training_records': len(records), 'independent_groups': len({groups[sid] for sid in ids}),
               'declared_family_count': len(confounding['families']),
               'all_families_single_independent_group': confounding['all_families_single_group'],
               'defined_binder_component_sum_available': sum(r['binder_component_sum'] is not None for r in records),
               'binder_total_conflict_records': sum(r['binder_total_conflict'] for r in records),
               'source_water_reported_in_L_m3': sum(r['water_source_unit'] == 'L/m3' for r in records),
               'water_binder_candidate_values': sum(r['candidate_ratios']['water_over_defined_reported_binder_components'] is not None for r in records),
               'models_trained': False, 'test_rows_used': 0, 'new_trusted_samples': 0,
               'inputs_unchanged': all(digest(p) == h for p, h in before.items())}
    if not summary['inputs_unchanged']:
        raise RuntimeError('组成核查输入改变')
    policy = {'model_training': False, 'automatic_feature_admission': False, 'zero_imputation': False,
              'scm_types_kept_separate': True, 'fiber_types_kept_separate': True,
              'family_labels_not_admitted_due_to_single_group_confounding': True,
              'water_volume_to_mass_conversion': None,
              'binder_total_conflicts_not_resolved_by_overwriting': True,
              'next_step': '先消除横线/未列组分的语义不确定及定义冲突，再预先约定可用特征敏感性分析；不按本轮误差删样本。'}
    output.mkdir(parents=True, exist_ok=False)
    for name, data in [('composition_records.json', records), ('field_coverage.json', coverage),
                       ('family_confounding.json', confounding), ('summary.json', summary),
                       ('feature_policy.json', policy), ('evidence_snapshot.json', config),
                       ('manifest.json', {'inputs_sha256': before, 'module_sha256': digest(ROOT / 'src/concrete_audit/composition_review.py'), 'script_sha256': digest(__file__)})]:
        (output / name).write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False) + '\n')
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = '128条组成诊断'
    ratio_names = list(records[0]['candidate_ratios'])
    append_excel_row(sheet, ['原始行', '来源', '原文声明家族', *ratio_names, '定义内组分和', '组分和单位/基准', '原文总量', '总量矛盾', '水量原单位', '人工状态'])
    for r in records:
        append_excel_row(sheet, [r['source_row'], r['paper'], r['declared_mix_family'], *r['candidate_ratios'].values(),
                                 r['binder_component_sum'], r['binder_sum_unit'], r['source_printed_binder_total'], r['binder_total_conflict'], r['water_source_unit'], '待核查'])
    sheet = book.create_sheet('组分原值与未知原因')
    append_excel_row(sheet, ['原始行', '来源', '组分', '原文值', '原单位/基准', '状态', '说明'])
    for r in records:
        for name, component in {**r['components'], **r['fibers']}.items():
            append_excel_row(sheet, [r['source_row'], r['paper'], name, component['value'], component['unit'], component['status'], r['source_notes']])
    sheet = book.create_sheet('家族与来源混杂')
    append_excel_row(sheet, ['原文声明家族', '训练记录数', '独立组数', '独立组ID'])
    for r in confounding['families']:
        append_excel_row(sheet, [r['declared_mix_family'], r['records'], r['group_count'], ','.join(r['independent_groups'])])
    sheet = book.create_sheet('原文证据')
    append_excel_row(sheet, ['原始行', '配方', '来源', '页码', '摘录', '说明'])
    for a in config['annotations']:
        append_excel_row(sheet, [a['source_row'], a['specimen_id'], a['paper'], ','.join(str(e['page']) for e in a['evidence']),
                                 '\n\n'.join(e['quote'] for e in a['evidence']), a['notes']])
    for sheet in book:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
    book.save(output / '胶材纤维与材料体系核查.xlsx')
    lines = ['# 胶凝材料、纤维和原文材料家族核查', '',
             '范围：现有严格训练128条、5个独立来源组；未按误差删样本，未重新训练或使用145条测试。全部比例为独立候选诊断，不自动授予可信资格或加入模型。', '',
             '## 原文声明家族与来源完全混杂', '',
             '|原文声明家族|训练记录|独立来源组|', '|---|---:|---:|']
    for r in confounding['families']:
        lines.append(f'|{r["declared_mix_family"]}|{r["records"]}|{r["group_count"]}|')
    lines += ['', '这些是原文自述的家族/用途标签，并非互斥的通用材料分类。本训练集每类仅1组；按来源留组时，每个验证记录的家族在该训练折都未见过。直接编码标签不能证明跨材料体系预测能力，也不能用强度阈值反推体系标签。', '',
              '## 组分证据覆盖', '', '|候选比例|有值/128|独立来源组|', '|---|---:|---:|']
    for r in coverage:
        lines.append(f'|{r["field"]}|{r["known_training_records"]}|{r["known_groups"]}|')
    lines += ['', 'FA、SF、GGBS及钢/PVA/PP纤维分别保留；未将未报告、原表横线或其他纤维类型都填0。纤维质量比不等于体积分数，也不表示不同纤维效果相同。', '',
              '## 胶材分母与水量单位', '',
              f'只有{summary["defined_binder_component_sum_available"]}条能对已定义且各项明确的原文胶材组分求完整和；其余不把缺项当0。',
              'Zhang四条：原文归一化胶材1+.06+.18+.19=1.43，按报告组分计算水胶比.223/1.43≈0.15594；不因此核实原库kg/m³基准。',
              'Al-Selmi六条：分项1202+301=1503，而原文表注1502；保存两个数值及冲突标记。605/1503≈0.40253只是按分项定义计算的候选，不擅自选定正确总量或直接用于训练。',
              'Yousri仅M-0.5-10FA-5SF一条在Cem/FA/SF三列都为明确正值，可对这一定义求和592；其他行的横线保持未知，不强行统一胶材总量。',
              f'Yousri11条、Brzozowski3条共{summary["source_water_reported_in_L_m3"]}条，原文Water列单位为L/m³。当前历史核心输入虽按质量使用，本模块不以数值相同证明kg/m³转换，不自动假定水密度；新的水胶比不跨单位计算。原Water及历史模型不回写。', '',
              '## 当前处理与下一步', '',
              '暂不把这些稀疏/有定义疑点的字段直接加进模型。先完成原表横线是否代表未掺的字段级核查、胶材定义与单位绑定，再确定预先声明的特征敏感性方案；家族标签保持诊断用途。当前来源数量不支持按每个家族各自建立可泛化模型。',
              '没有通过删除UHPC/ECC记录来追求更高R²，未做目标驱动的材料分类，也不再扩大参数搜索。', '',
              '关键函数：complete_component_sum对任何未知组分返回未知；derive_composition分开保存分项和、原文总量与候选比例；family_confounding检查家族在各来源折中的可见性。', '',
              '输出：composition_records.json、field_coverage.json、family_confounding.json、feature_policy.json及独立Excel；证据均重新与实际PDF页码/摘录和版本校验。']
    (output / 'report.md').write_text('\n'.join(lines) + '\n')
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    run(parser.parse_args().output)
