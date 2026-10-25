"""第一、二步：只读盘点、字段画像和已有处理历史线索，不执行文献自动授信。"""
import argparse
from collections import Counter
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
from urllib.parse import unquote
import openpyxl


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def profile(values):
    """分别计数空白、占位、零和公式；零值不等价于真实未添加。"""
    values=list(values)
    return {'cells':len(values), 'blank':sum(v is None or v=='' for v in values),
            'placeholder_text':sum(isinstance(v,str) and v.strip() in ('-','/') for v in values),
            'numeric_zero':sum(isinstance(v,(int,float)) and not isinstance(v,bool) and v==0 for v in values),
            'formula':sum(isinstance(v,str) and v.startswith('=') for v in values)}


def run(attachments, output):
    files=sorted(p for p in attachments.rglob('*') if p.is_file())
    before={str(p):digest(p) for p in files}
    output.mkdir(parents=True,exist_ok=False)
    inventory=[{'path':str(p),'display_name':unquote(p.name),'bytes':p.stat().st_size,
                'sha256':before[str(p)],'suffix':p.suffix.lower()} for p in files]
    workbooks=[];profiles=[];fields=[]
    for p in files:
        if p.suffix.lower()!='.xlsx':continue
        w=openpyxl.load_workbook(p,data_only=False)
        cached=openpyxl.load_workbook(p,data_only=True,read_only=True)
        entries=[]
        for s in w:
            first=7 if s.title=='Data' else 2
            matrix=list(s.iter_rows(min_row=first,values_only=True))
            nonempty=[r for r in matrix if any(v is not None for v in r)]
            field_map={}
            if s.title=='Data' and 'Parameter Explanation' in w:
                for r in w['Parameter Explanation'].iter_rows(min_row=2,values_only=True):
                    if isinstance(r[0],int) and r[3]:field_map[r[0]+1]=r[3]
                    if isinstance(r[0],int) and r[3]:fields.append({'file':str(p),'number':r[0],
                        'symbol':r[3],'unit':r[4],'definition':r[6]})
            elif s.max_row:
                field_map={i+1:str(v) for i,v in enumerate(next(s.values,())) if v is not None}
            is_samples=s.title in ('Data','初筛候选527','待核查清单')
            duplicates=sum(n-1 for n in Counter(nonempty).values() if n>1) if is_samples else None
            formula_coords=[c.coordinate for row in s for c in row if c.data_type=='f']
            cs=cached[s.title]
            cache_missing=[]
            if formula_coords:
                wanted=set(formula_coords)
                for ri,row in enumerate(cs,1):
                    for ci,c in enumerate(row,1):
                        coord=f'{openpyxl.utils.get_column_letter(ci)}{ri}'
                        if coord in wanted and c.value is None:cache_missing.append(coord)
            entries.append({'sheet':s.title,'rows':s.max_row,'columns':s.max_column,
                'data_start_row':first,'nonempty_body_rows':len(nonempty),
                'sample_count':len(nonempty) if is_samples else None,
                'duplicate_excess_rows_raw_cells':duplicates,
                'formula_cells':len(formula_coords),'formula_missing_cache':len(cache_missing),
                'formula_missing_cache_examples':cache_missing[:15],
                'comments':sum(c.comment is not None for row in s for c in row),
                'excel_errors':[{ 'cell':c.coordinate,'value':c.value} for row in s for c in row if c.data_type=='e']})
            for col in range(1,s.max_column+1):
                profiles.append({'file':str(p),'sheet':s.title,'column':col,
                    'field':field_map.get(col,f'列{col}'),**profile(r[col-1] for r in matrix)})
        workbooks.append({'path':str(p),'sheets':entries})
        w.close();cached.close()
    assert all(digest(p)==before[str(p)] for p in files),'输入发生变化'
    for name,obj in [('files.json',inventory),('workbooks.json',workbooks),
                     ('field_profiles.json',profiles),('field_dictionary.json',fields)]:
        (output/name).write_text(json.dumps(obj,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    manifest={'created_at_utc':datetime.now(timezone.utc).isoformat(),'inputs_unchanged':True,
              'counts_by_extension':dict(Counter(x['suffix'] for x in inventory)),
              'script_sha256':digest(Path(__file__)),'python_stage':'仅文件盘点和处理历史梳理',
              'model_training':False,'external_llm_called':False}
    (output/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))
    print(json.dumps(manifest,ensure_ascii=False,indent=2))
    for book in workbooks:
        print(Path(book['path']).name)
        print([{k:s[k] for k in ('sheet','sample_count','duplicate_excess_rows_raw_cells','formula_cells','formula_missing_cache')}
               for s in book['sheets'] if s['sample_count'] is not None])

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--attachments',type=Path,default=Path('/workspace/attachments'))
    parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args();run(args.attachments,args.output)
